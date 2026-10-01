"""Workbook layout of the LUDA I/F template (seed asset ``if-excel``) and its reader/writer."""

import io
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

IF_SHEET = "인터페이스 리스트"
SYSTEM_SHEET = "시스템 연동정보"
IF_COLUMNS = (
    ("if_code", "I/F ID"),
    ("name", "I/F 명"),
    ("source", "송신 시스템"),
    ("target", "수신 시스템"),
    ("link_type", "연동 방식"),
    ("schedule", "주기"),
    ("description", "데이터 항목"),
    ("daily_volume", "건수(일)"),
    ("owner", "담당자"),
    ("status", "상태"),
    ("notes", "비고"),
)
SYSTEM_COLUMNS = (
    ("name", "시스템 명"),
    ("short_name", "약칭"),
    ("type", "유형"),
    ("owner_dept", "담당 부서"),
    ("hosting", "운영 형태"),
    ("db_type", "DB 종류"),
    ("connection", "접속 정보(마스킹)"),
    ("notes", "비고"),
)
IF_REQUIRED = ("if_code", "name", "source", "target")
MAX_ROWS = 5000

LINK_TYPE_LABELS = {
    "db_link": "DB Link",
    "api": "API",
    "file": "File",
    "mq": "MQ",
    "eai": "EAI",
    "other": "기타",
}
STATUS_LABELS = {"planned": "계획", "developing": "개발", "operating": "운영", "retired": "폐기"}
HOSTING_LABELS = {"on_premise": "온프레미스", "cloud": "클라우드"}

_LINK_ALIASES = {
    "dblink": "db_link",
    "db": "db_link",
    "dbtodb": "db_link",
    "api": "api",
    "rest": "api",
    "restapi": "api",
    "soap": "api",
    "webservice": "api",
    "file": "file",
    "파일": "file",
    "ftp": "file",
    "sftp": "file",
    "mq": "mq",
    "kafka": "mq",
    "eai": "eai",
    "esb": "eai",
    "other": "other",
    "기타": "other",
}
_STATUS_ALIASES = {
    "planned": "planned",
    "계획": "planned",
    "예정": "planned",
    "developing": "developing",
    "개발": "developing",
    "개발중": "developing",
    "operating": "operating",
    "운영": "operating",
    "운영중": "operating",
    "retired": "retired",
    "폐기": "retired",
    "중지": "retired",
}
_HOSTING_ALIASES = {
    "on_premise": "on_premise",
    "onpremise": "on_premise",
    "온프레미스": "on_premise",
    "온프렘": "on_premise",
    "cloud": "cloud",
    "클라우드": "cloud",
}
_SYSTEM_TYPES = ("MES", "ERP", "LIMS", "WMS", "SCADA", "GROUPWARE", "OTHER")
_SYSTEM_TYPE_ALIASES = {"그룹웨어": "GROUPWARE", "기타": "OTHER"}


class WorkbookError(Exception):
    def __init__(self, code: str, detail: dict[str, Any] | None = None) -> None:
        super().__init__(code)
        self.code = code
        self.detail = detail


@dataclass
class SheetRow:
    row: int
    values: dict[str, str | None] = field(default_factory=dict)


@dataclass
class ParsedWorkbook:
    interfaces: list[SheetRow]
    systems: list[SheetRow]


def _key(value: str) -> str:
    return "".join(value.lower().split()).replace("-", "").replace("_", "")


def normalize_link_type(value: str | None) -> str | None:
    if value is None:
        return "other"
    return _LINK_ALIASES.get(_key(value))


def normalize_status(value: str | None) -> str | None:
    if value is None:
        return "operating"
    return _STATUS_ALIASES.get(_key(value))


def normalize_hosting(value: str | None) -> str | None:
    return None if value is None else _HOSTING_ALIASES.get(_key(value))


def normalize_system_type(value: str | None) -> str:
    if value is None:
        return "OTHER"
    upper = value.strip().upper()
    if upper in _SYSTEM_TYPES:
        return upper
    return _SYSTEM_TYPE_ALIASES.get(value.strip(), "OTHER")


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _read_sheet(rows: list[tuple[object, ...]], columns: tuple[tuple[str, str], ...], sheet: str) -> list[SheetRow]:
    if not rows:
        raise WorkbookError("invalid_workbook", {"sheet": sheet, "missing_columns": [h for _, h in columns]})
    header = [_cell_text(v) or "" for v in rows[0]]
    positions: dict[str, int] = {}
    missing: list[str] = []
    for key, label in columns:
        if label in header:
            positions[key] = header.index(label)
        else:
            missing.append(label)
    if missing:
        raise WorkbookError("invalid_workbook", {"sheet": sheet, "missing_columns": missing})
    out: list[SheetRow] = []
    for idx, raw in enumerate(rows[1:], start=2):
        values = {k: _cell_text(raw[p]) if p < len(raw) else None for k, p in positions.items()}
        if any(v is not None for v in values.values()):
            out.append(SheetRow(row=idx, values=values))
        if len(out) > MAX_ROWS:
            raise WorkbookError("too_many_rows", {"sheet": sheet, "max": MAX_ROWS})
    return out


def parse_workbook(data: bytes) -> ParsedWorkbook:
    """Read values only (no formulas or macros are evaluated)."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise WorkbookError("invalid_workbook") from exc
    try:
        if IF_SHEET not in wb.sheetnames:
            raise WorkbookError("invalid_workbook", {"missing_sheet": IF_SHEET})
        interfaces = _read_sheet(list(wb[IF_SHEET].iter_rows(values_only=True)), IF_COLUMNS, IF_SHEET)
        systems: list[SheetRow] = []
        if SYSTEM_SHEET in wb.sheetnames:
            systems = _read_sheet(list(wb[SYSTEM_SHEET].iter_rows(values_only=True)), SYSTEM_COLUMNS, SYSTEM_SHEET)
        return ParsedWorkbook(interfaces=interfaces, systems=systems)
    finally:
        wb.close()


_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]], first: bool) -> None:
    ws = wb.active if first and wb.active is not None else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for row in rows:
        ws.append([None if v is None else v for v in row])
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) * 2 + 4)
    ws.freeze_panes = "A2"


def build_workbook(interface_rows: list[list[Any]], system_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    _write_sheet(wb, IF_SHEET, [h for _, h in IF_COLUMNS], interface_rows, first=True)
    _write_sheet(wb, SYSTEM_SHEET, [h for _, h in SYSTEM_COLUMNS], system_rows, first=False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
