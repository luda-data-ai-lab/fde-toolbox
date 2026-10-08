import io
import json
import sqlite3
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from app.modules.exmigrate.analyzer import analyze_workbook
from app.modules.exmigrate.erd import draft_erd, erd_issues, to_identifier, to_mermaid
from app.modules.exmigrate.schemas import Erd
from app.modules.exmigrate.scripts import ddl, load_order
from tests.conftest import XLSX, World


def _workbook() -> bytes:
    wb = Workbook()
    products = wb.active
    assert products is not None
    products.title = "제품"
    products.append(["제품 마스터"])
    products.append(["제품코드", "제품명", "단가"])
    products.append(["P-1", "도료 A", 1200.5])
    products.append(["P-2", "도료 B", 900])
    orders = wb.create_sheet("Order List")
    orders.append(["order_no", "제품코드", "수량", "금액", "제품명", "비고"])
    orders.append([1, "P-1", 3, "=C2*VLOOKUP(B2,제품!$A$3:$C$10,3,FALSE)", "=VLOOKUP(B2,'제품'!A:C,2,FALSE)", None])
    orders.append([None, None, None, None, None, None])
    orders.append([2, "P-2", 5, "=SUM(C2:C4)", '=IFERROR(INDEX(제품!B:B,MATCH(B4,제품!A:A,0)),"")', "x"])
    orders.append([3, "P-1", "많음", "=LAMBDA(x,x)(1)", None, None])
    orders.merge_cells("F2:F3")
    wb.create_sheet("빈 시트")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _base(world: World) -> str:
    return f"/api/v1/t/{world.a.tenant_id}/exmigrate"


def _upload(world: World, data: bytes | None = None, name: str = "orders.xlsx") -> dict[str, Any]:
    r = world.a.fde.post(
        f"{_base(world)}/analyses",
        data={"engagement_id": world.a.ids["engagement_id"]},
        files={"file": (name, data or _workbook(), XLSX)},
    )
    assert r.status_code == 201, r.text
    body: dict[str, Any] = r.json()
    return body


def _actions(world: World) -> list[str]:
    logs = world.a.fde.get(f"/api/v1/t/{world.a.tenant_id}/audit-logs", params={"limit": 200}).json()["items"]
    return [x["action"] for x in logs]


def test_structure_analysis() -> None:
    report = analyze_workbook(_workbook())
    product, order, empty = report.sheets
    assert product.header_row == 2 and product.data_rows == 2
    assert [(c.name, c.inferred_type) for c in product.columns] == [
        ("제품코드", "text"),
        ("제품명", "text"),
        ("단가", "decimal"),
    ]
    assert order.header_row == 1 and order.data_rows == 4
    cols = {c.name: c for c in order.columns}
    assert cols["order_no"].inferred_type == "integer" and cols["order_no"].null_ratio == 0.25
    assert cols["수량"].inferred_type == "text" and cols["제품코드"].distinct == 2
    codes = {w.code for w in order.warnings}
    assert {"merged_cells", "blank_rows", "mixed_types"} <= codes
    assert next(w for w in order.warnings if w.code == "blank_rows").where == "3"
    assert [w.code for w in empty.warnings] == ["empty_sheet"] and not empty.columns


def test_formula_analysis() -> None:
    report = analyze_workbook(_workbook())
    order = report.sheets[1].formulas
    assert order.count == 5 and report.formula_count == 5
    assert order.cells[0].cell == "D2"
    funcs = {f.name: f for f in order.functions}
    assert funcs["VLOOKUP"].count == 2 and funcs["VLOOKUP"].simple
    assert not funcs["LAMBDA"].simple and order.complex == 1
    assert [(r.sheet, r.count) for r in order.references] == [("제품", 4)]
    assert {(lk.column, lk.target_sheet, lk.target_column) for lk in order.lookups} == {("B", "제품", "A")}
    assert not report.has_macros


def test_invalid_workbook_rejected(world: World) -> None:
    r = world.a.fde.post(
        f"{_base(world)}/analyses",
        data={"engagement_id": world.a.ids["engagement_id"]},
        files={"file": ("bad.xlsx", b"not a zip", XLSX)},
    )
    assert r.status_code == 400 and r.json()["error"]["code"] == "workbook_invalid"
    r = world.a.fde.post(
        f"{_base(world)}/analyses",
        data={"engagement_id": world.a.ids["engagement_id"]},
        files={"file": ("x.csv", b"a,b", "text/csv")},
    )
    assert r.status_code == 400


def test_erd_draft() -> None:
    erd = draft_erd(analyze_workbook(_workbook()))
    assert [t.name for t in erd.tables] == ["제품", "order_list"]
    product, order = erd.tables
    assert [c.name for c in product.columns if c.primary_key] == ["제품코드"]
    assert not order.columns[0].primary_key  # order_no has a blank row
    assert [(r.from_table, r.from_column, r.to_table, r.to_column, r.origin) for r in erd.relations] == [
        ("order_list", "제품코드", "제품", "제품코드", "formula")
    ]
    assert erd_issues(erd) == []
    assert to_identifier("1st Qty (kg)", "x") == "c_1st_qty_kg"
    mmd = to_mermaid(erd)
    assert 't1["제품"] {' in mmd and 'text col1 PK "제품코드"' in mmd and 't1 ||--o{ order_list : "제품코드"' in mmd


def test_erd_issues() -> None:
    erd = Erd.model_validate(
        {
            "tables": [
                {"name": "a", "columns": [{"name": "x"}, {"name": "X"}]},
                {"name": "A", "columns": [{"name": "1bad"}]},
            ],
            "relations": [{"from_table": "a", "from_column": "nope", "to_table": "zz", "to_column": "x"}],
        }
    )
    assert erd_issues(erd) == [
        "duplicate_column:a.X",
        "duplicate_table:A",
        "invalid_name:A.1bad",
        "unknown_relation_source:a.nope",
        "unknown_relation_target:zz.x",
    ]


def test_ddl_dialects() -> None:
    erd = draft_erd(analyze_workbook(_workbook()))
    pg = ddl(erd, "postgresql", "orders.xlsx")
    assert 'CREATE TABLE "제품"' in pg and '"단가" NUMERIC' in pg and '"order_no" BIGINT' in pg
    assert 'FOREIGN KEY ("제품코드") REFERENCES "제품" ("제품코드");' in pg
    ms = ddl(erd, "mssql", "orders.xlsx")
    assert "[제품코드] NVARCHAR(450) NOT NULL" in ms and "[비고] NVARCHAR(MAX)" in ms
    assert "ALTER TABLE [order_list] ADD CONSTRAINT [fk_order_list_1]" in ms
    lite = ddl(erd, "sqlite", "a\nb.xlsx")
    assert lite.startswith("-- Generated by FDE Toolbox ExMigrate from a b.xlsx (sqlite)")
    conn = sqlite3.connect(":memory:")
    conn.executescript(lite)
    assert [t.name for t in load_order(erd)] == ["제품", "order_list"]


def test_upload_erd_and_exports(world: World, tmp_path: Path) -> None:
    data = _workbook()
    a = _upload(world, data)
    assert a["report"]["formula_count"] == 5 and a["sheets"] == 3 and not a["erd_confirmed"]
    listed = world.a.fde.get(f"{_base(world)}/analyses", params={"engagement_id": world.a.ids["engagement_id"]})
    assert a["id"] in [x["id"] for x in listed.json()["items"]]
    url = f"{_base(world)}/analyses/{a['id']}"
    assert world.a.fde.get(f"{url}/ddl.sql").status_code == 409
    erd = world.a.fde.get(f"{url}/erd").json()
    assert erd["confirmed"] is False and erd["issues"] == []

    bad = json.loads(json.dumps(erd["erd"]))
    bad["tables"][0]["columns"].append({"name": bad["tables"][0]["columns"][0]["name"]})
    assert world.a.fde.put(f"{url}/erd", json=bad).json()["issues"]
    r = world.a.fde.post(f"{url}/erd/confirm")
    assert r.status_code == 422 and r.json()["error"]["code"] == "erd_invalid"

    good = erd["erd"]
    good["tables"][1]["columns"][0]["primary_key"] = True
    assert world.a.fde.put(f"{url}/erd", json=good).status_code == 200
    confirmed = world.a.fde.post(f"{url}/erd/confirm").json()
    assert confirmed["confirmed"] and confirmed["confirmed_at"]
    assert world.a.fde.get(url).json()["erd_confirmed"]

    report = world.a.fde.get(f"{url}/report.md")
    assert report.status_code == 200 and "## Sheet: Order List" in report.text
    assert "| VLOOKUP | 2 | yes |" in report.text
    assert "erDiagram" in world.a.fde.get(f"{url}/erd.mmd").text
    assert 'CONSTRAINT "pk_order_list" PRIMARY KEY ("order_no")' in world.a.fde.get(f"{url}/ddl.sql").text

    z = world.a.fde.get(f"{url}/scripts.zip", params={"dialect": "sqlite"})
    assert z.status_code == 200
    with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
        assert sorted(zf.namelist()) == ["README.md", "load.py", "mapping.json", "schema.sql"]
        zf.extractall(tmp_path)
    db = tmp_path / "target.db"
    conn = sqlite3.connect(db)
    conn.executescript((tmp_path / "schema.sql").read_text(encoding="utf-8"))
    conn.close()
    (tmp_path / "orders.xlsx").write_bytes(data)
    run = subprocess.run(
        [sys.executable, "load.py", "orders.xlsx", "--dialect", "sqlite", "--url", str(db)],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        check=False,
    )
    assert run.returncode == 0, run.stderr
    conn = sqlite3.connect(db)
    assert conn.execute('SELECT "제품코드", "단가" FROM "제품" ORDER BY 1').fetchall() == [
        ("P-1", 1200.5),
        ("P-2", 900),
    ]
    assert conn.execute('SELECT count(*) FROM "order_list"').fetchone() == (3,)
    assert {
        "exmigrate.analysis_create",
        "exmigrate.erd_update",
        "exmigrate.erd_confirm",
        "exmigrate.export",
    } <= set(_actions(world))

    put = world.a.fde.put(f"{url}/erd", json=good).json()
    assert put["confirmed"] is False


def test_permissions_and_scope(world: World) -> None:
    a = _upload(world)
    url = f"{_base(world)}/analyses/{a['id']}"
    assert world.a.client_user.post(f"{url}/erd/confirm").status_code == 403
    assert world.a.client_user.get(f"{url}/scripts.zip").status_code == 403
    assert world.a.client_user.get(url).status_code == 200
    r = world.a.fde.post(
        f"{_base(world)}/analyses",
        data={"engagement_id": world.b.ids["engagement_id"]},
        files={"file": ("o.xlsx", _workbook(), XLSX)},
    )
    assert r.status_code in (404, 422), r.text
    assert world.a.fde.delete(url).status_code == 204
    assert world.a.fde.get(url).status_code == 404
