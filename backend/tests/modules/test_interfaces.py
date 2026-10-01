import csv
import io
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.modules.interfaces import excel
from tests.conftest import XLSX, World

SEED = Path(__file__).resolve().parents[2] / "seeds" / "assets" / "if-excel-template.json"


def _upload(world: World, data: bytes, name: str = "if.xlsx") -> dict[str, object]:
    base = f"/api/v1/t/{world.a.tenant_id}/interfaces"
    r = world.a.fde.post(f"{base}/uploads", files={"file": (name, data, XLSX)})
    assert r.status_code == 201, r.text
    body: dict[str, object] = r.json()
    return body


def test_template_matches_seed_asset(world: World) -> None:
    payload = json.loads(SEED.read_text(encoding="utf-8"))["assets"][0]["payload"]
    sheets = {s["name"]: s["columns"] for s in payload["sheets"]}
    assert sheets == {
        excel.IF_SHEET: [h for _, h in excel.IF_COLUMNS],
        excel.SYSTEM_SHEET: [h for _, h in excel.SYSTEM_COLUMNS],
    }
    r = world.a.client_user.get(f"/api/v1/t/{world.a.tenant_id}/interfaces/template.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [excel.IF_SHEET, excel.SYSTEM_SHEET]
    assert [c.value for c in wb[excel.IF_SHEET][1]] == sheets[excel.IF_SHEET]


def test_upload_validation_and_apply(world: World) -> None:
    base = f"/api/v1/t/{world.a.tenant_id}/interfaces"
    rows = [
        ["IF-TA-001", "기존 갱신", "ERP TA", "ERP TA", "REST API", "실시간", "주문", "1,200", "김", "운영", None],
        ["IF-NEW-01", "실적", "erp ta", "MES 신규", "DB Link", "1시간", None, None, None, None, None],
        ["IF-NEW-02", "검사", "LIMS", "MES 신규", "파일", None, None, None, None, "개발중", None],
        ["IF-NEW-02", "중복", "ERP TA", "ERP TA", "API", None, None, None, None, None, None],
        [None, "코드 없음", "ERP TA", None, "smoke signal", None, None, "many", None, "unknown", None],
    ]
    systems = [["MES 신규", "MES", "MES", "생산팀", "온프레미스", "MSSQL", "***", None]]
    up = _upload(world, excel.build_workbook(rows, systems))
    result = up["result"]
    assert isinstance(result, dict)
    assert result["summary"] == {"total": 5, "valid": 3, "invalid": 2, "create": 2, "update": 1}
    by_row = {r["row"]: r for r in result["rows"]}
    assert by_row[2]["values"]["daily_volume"] == 1200 and by_row[2]["values"]["link_type"] == "api"
    assert by_row[4]["values"]["status"] == "developing"
    assert {"field": "if_code", "code": "duplicate_in_file"} in by_row[5]["errors"]
    codes = {(e["field"], e["code"]) for e in by_row[6]["errors"]}
    assert codes == {
        ("if_code", "required"),
        ("target", "required"),
        ("link_type", "unknown_link_type"),
        ("status", "unknown_status"),
        ("daily_volume", "invalid_number"),
    }
    unreg = {u["name"]: u for u in result["unregistered_systems"]}
    assert set(unreg) == {"MES 신규", "LIMS"}
    assert unreg["MES 신규"]["in_system_sheet"] and unreg["MES 신규"]["rows"] == 2
    assert unreg["MES 신규"]["system"]["hosting"] == "on_premise"

    assert world.a.client_user.post(f"{base}/uploads/{up['id']}/apply", json={}).status_code == 403
    applied = world.a.fde.post(f"{base}/uploads/{up['id']}/apply", json={"register_systems": ["MES 신규"]}).json()
    assert applied["status"] == "applied"
    summary = applied["result"]["applied"]
    assert summary["created"] == 1 and summary["updated"] == 1 and summary["systems_created"] == ["MES 신규"]
    assert [s["if_code"] for s in summary["skipped"]] == ["IF-NEW-02"]
    again = world.a.fde.post(f"{base}/uploads/{up['id']}/apply", json={})
    assert again.status_code == 409 and again.json()["error"]["code"] == "upload_already_applied"

    listed = world.a.fde.get(base, params={"q": "if-"}).json()["items"]
    assert sorted(i["if_code"] for i in listed) == ["IF-NEW-01", "IF-TA-001"]
    updated = next(i for i in listed if i["if_code"] == "IF-TA-001")
    assert updated["name"] == "기존 갱신" and updated["link_type"] == "api"
    sys_names = {s["name"]: s for s in world.a.fde.get(f"/api/v1/t/{world.a.tenant_id}/systems").json()["items"]}
    assert sys_names["MES 신규"]["db_type"] == "MSSQL"
    assert sys_names["MES 신규"]["tenant_id"] == world.a.tenant_id
    upload_file = world.a.fde.get(f"/api/v1/t/{world.a.tenant_id}/files/{up['file_id']}").json()
    assert upload_file["owner_type"] == "interface_upload" and upload_file["owner_id"] == up["id"]


def test_bad_workbooks_rejected(world: World) -> None:
    base = f"/api/v1/t/{world.a.tenant_id}/interfaces"
    r = world.a.fde.post(f"{base}/uploads", files={"file": ("x.xlsx", b"not a zip", XLSX)})
    assert r.status_code == 400 and r.json()["error"]["code"] == "invalid_workbook"
    wb = Workbook()
    assert wb.active is not None
    wb.active.title = "Sheet"
    buf = io.BytesIO()
    wb.save(buf)
    r = world.a.fde.post(f"{base}/uploads", files={"file": ("x.xlsx", buf.getvalue(), XLSX)})
    assert r.json()["error"]["detail"] == {"missing_sheet": excel.IF_SHEET}
    wb.active.title = excel.IF_SHEET
    wb.active.append(["I/F ID", "I/F 명"])
    buf = io.BytesIO()
    wb.save(buf)
    r = world.a.fde.post(f"{base}/uploads", files={"file": ("x.xlsx", buf.getvalue(), XLSX)})
    assert r.status_code == 400 and "송신 시스템" in r.json()["error"]["detail"]["missing_columns"]
    r = world.a.fde.post(f"{base}/uploads", files={"file": ("x.csv", b"a,b", "text/csv")})
    assert r.json()["error"]["code"] == "file_type_not_allowed"
    assert world.a.client_admin.post(f"{base}/uploads", files={"file": ("x.xlsx", b"x", XLSX)}).status_code == 403


def test_crud_dashboard_graph_and_exports(world: World) -> None:
    base = f"/api/v1/t/{world.a.tenant_id}/interfaces"
    fde = world.a.fde
    erp = world.a.ids["system_id"]
    mes = fde.post(f"/api/v1/t/{world.a.tenant_id}/systems", json={"name": "MES", "type": "MES"}).json()["id"]
    body = {"if_code": "IF-2", "name": '=HYPERLINK("x")', "source_system_id": erp, "target_system_id": mes}
    created = fde.post(base, json={**body, "link_type": "mq", "status": "planned", "tenant_id": world.b.tenant_id})
    assert created.status_code == 201 and created.json()["tenant_id"] == world.a.tenant_id
    iid = created.json()["id"]
    assert fde.post(base, json=body).json()["error"]["code"] == "if_code_taken"
    foreign = {**body, "if_code": "IF-3", "target_system_id": world.b.ids["system_id"]}
    assert fde.post(base, json=foreign).json()["error"]["code"] == "invalid_reference"
    assert fde.patch(f"{base}/{iid}", json={"if_code": "IF-TA-001"}).status_code == 409
    assert fde.patch(f"{base}/{iid}", json={"name": None}).status_code == 422
    assert fde.patch(f"{base}/{iid}", json={"status": "operating"}).json()["status"] == "operating"
    assert world.a.client_user.patch(f"{base}/{iid}", json={"status": "retired"}).status_code == 403

    assert [i["id"] for i in fde.get(base, params={"system_id": mes}).json()["items"]] == [iid]
    assert [i["id"] for i in fde.get(base, params={"link_type": "mq"}).json()["items"]] == [iid]

    dash = fde.get(f"{base}/dashboard").json()
    assert dash["total"] == 2 and dash["by_link_type"] == {"mq": 1, "other": 1}
    counts = {c["name"]: (c["outgoing"], c["incoming"]) for c in dash["by_system"]}
    assert counts == {"ERP TA": (2, 1), "MES": (0, 1)}

    g = fde.get(f"{base}/graph").json()
    assert {n["name"]: n["degree"] for n in g["nodes"]} == {"ERP TA": 3, "MES": 1}
    assert {(e["source"], e["target"]) for e in g["edges"]} == {(erp, erp), (erp, mes)}
    assert len(fde.get(f"{base}/graph", params={"link_type": "mq"}).json()["edges"]) == 1

    r = world.a.client_admin.get(f"{base}/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb[excel.IF_SHEET]
    assert [c.value for c in ws[2]][:5] == ["IF-2", '=HYPERLINK("x")', "ERP TA", "MES", "MQ"]
    assert ws.cell(row=2, column=2).data_type == "s"
    assert wb[excel.SYSTEM_SHEET].cell(row=2, column=1).value == "ERP TA"
    reparsed = excel.parse_workbook(r.content)
    assert [x.values["if_code"] for x in reparsed.interfaces] == ["IF-2", "IF-TA-001"]

    r = world.a.client_admin.get(f"{base}/export.csv")
    rows = list(csv.reader(io.StringIO(r.text.lstrip("\ufeff"))))
    assert rows[0][0] == "I/F ID" and rows[1][1] == '\'=HYPERLINK("x")'
    assert world.a.client_user.get(f"{base}/export.csv").status_code == 403
    actions = [x["action"] for x in fde.get(f"/api/v1/t/{world.a.tenant_id}/audit-logs").json()["items"]]
    assert actions.count("interface.export") == 2 and "interface.create" in actions

    in_use = fde.delete(f"/api/v1/t/{world.a.tenant_id}/systems/{mes}")
    assert in_use.status_code == 409 and in_use.json()["error"]["code"] == "conflict"
    assert fde.delete(f"{base}/{iid}").status_code == 204
    assert fde.delete(f"/api/v1/t/{world.a.tenant_id}/systems/{mes}").status_code == 204
