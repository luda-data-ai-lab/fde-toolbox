"""Fictional MES sample tenant (한빛정밀, CNC-machined auto parts) filling every module for demos."""

import io
from datetime import date

from fastapi import UploadFile
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.assets.models import AssetItem
from app.core.assets.service import latest
from app.core.schemas import AssetRef
from app.core.systems.models import System
from app.core.systems.schemas import SystemIn
from app.core.systems.service import create_system
from app.core.tenancy.context import Principal, TenantContext
from app.core.tenants.models import Tenant
from app.core.tenants.schemas import EngagementIn, TenantIn
from app.core.tenants.service import create_engagement, create_tenant
from app.core.users.schemas import UserIn
from app.core.users.service import create_user
from app.db.session import tenant_scope
from app.modules.agenthub.schemas import InstanceIn
from app.modules.agenthub.service import create_instance
from app.modules.devtracker.schemas import IssueIn, ProjectIn, TaskIn
from app.modules.devtracker.service import create_issue, create_project, create_task
from app.modules.discoveryq.schemas import ActionItemIn, InsightIn, SessionIn, SessionQuestionIn, SubjectIn
from app.modules.discoveryq.service import (
    add_question,
    create_action_item,
    create_insight,
    create_session,
    create_subject,
)
from app.modules.exmigrate.service import confirm_erd, create_analysis, draft_of
from app.modules.flowdesk.schemas import FlowGraph, FlowIn, GraphEdge, GraphNode, PairIn, Position, SnapshotIn
from app.modules.flowdesk.service import create_flow, create_snapshot, pair_flow
from app.modules.interfaces.excel import build_workbook
from app.modules.interfaces.schemas import ApplyIn
from app.modules.interfaces.service import apply_upload, upload_workbook
from app.modules.ontomap.candidates import run_extract
from app.modules.ontomap.concepts import create_attribute, create_concept, create_relation
from app.modules.ontomap.mappings import create_mapping
from app.modules.ontomap.schemas import (
    AliasIn,
    AttributeIn,
    Cardinality,
    ConceptIn,
    DataType,
    ExtractIn,
    MappingIn,
    RelationIn,
    TermIn,
    UpperRef,
)
from app.modules.ontomap.service import create_term
from app.modules.specforge.schemas import DocumentIn, DocumentPatch, SpecSources
from app.modules.specforge.service import create_document, update_document

CODE = "HANBIT"
FDE_EMAIL = "fde@hanbit.local"
CLIENT_EMAIL = "client@hanbit.local"

SYSTEMS = [
    SystemIn(
        name="생산관리시스템",
        short_name="MES",
        type="MES",
        owner_dept="생산관리팀",
        hosting="on_premise",
        db_type="MSSQL",
    ),
    SystemIn(
        name="전사자원관리",
        short_name="ERP",
        type="ERP",
        owner_dept="경영지원팀",
        hosting="on_premise",
        db_type="Oracle",
    ),
    SystemIn(
        name="품질관리시스템",
        short_name="QMS",
        type="OTHER",
        owner_dept="품질보증팀",
        hosting="on_premise",
        db_type="MSSQL",
    ),
    SystemIn(
        name="제품수명주기관리",
        short_name="PLM",
        type="OTHER",
        owner_dept="생산기술팀",
        hosting="cloud",
        db_type="PostgreSQL",
    ),
    SystemIn(
        name="설비데이터수집",
        short_name="SCADA",
        type="SCADA",
        owner_dept="설비보전팀",
        hosting="on_premise",
        db_type="Historian",
    ),
    SystemIn(
        name="창고관리시스템", short_name="WMS", type="WMS", owner_dept="물류팀", hosting="cloud", db_type="PostgreSQL"
    ),
]

INTERFACES = [
    [
        "IF-HB-001",
        "작업지시 수신",
        "ERP",
        "MES",
        "DB Link",
        "1시간",
        "작업지시번호, 품목코드, 지시수량, 납기",
        120,
        "한생산",
        "운영",
        None,
    ],
    [
        "IF-HB-002",
        "생산실적 송신",
        "MES",
        "ERP",
        "DB Link",
        "1시간",
        "실적번호, 양품수량, 불량수량",
        900,
        "한생산",
        "운영",
        None,
    ],
    [
        "IF-HB-003",
        "BOM·라우팅 동기화",
        "PLM",
        "MES",
        "API",
        "일 1회",
        "품목코드, 공정순서, 표준 사이클타임",
        30,
        "오기술",
        "운영",
        None,
    ],
    [
        "IF-HB-004",
        "설비 가동 상태 수집",
        "SCADA",
        "MES",
        "MQ",
        "실시간",
        "설비코드, 가동/비가동, 알람코드",
        150000,
        "강설비",
        "운영",
        None,
    ],
    [
        "IF-HB-005",
        "가공 치수 측정값",
        "MES",
        "QMS",
        "API",
        "실시간",
        "LOT번호, 측정항목, 측정값",
        6000,
        "윤품질",
        "운영",
        None,
    ],
    [
        "IF-HB-006",
        "SPC 판정 결과",
        "QMS",
        "MES",
        "API",
        "실시간",
        "LOT번호, 판정, 관리한계 이탈",
        6000,
        "윤품질",
        "운영",
        None,
    ],
    [
        "IF-HB-007",
        "부적합 등록",
        "MES",
        "QMS",
        "API",
        "실시간",
        "부적합번호, 불량유형, 수량",
        40,
        "윤품질",
        "개발",
        "2차 오픈",
    ],
    [
        "IF-HB-008",
        "완제품 입고",
        "MES",
        "WMS",
        "DB Link",
        "30분",
        "LOT번호, 품목코드, 수량",
        300,
        "서물류",
        "운영",
        None,
    ],
    ["IF-HB-009", "출하 지시", "ERP", "WMS", "API", "30분", "출하번호, 고객사, 품목, 수량", 80, "서물류", "운영", None],
    [
        "IF-HB-010",
        "공구 교체 이력",
        "MES",
        "PLM",
        "File",
        "일 1회",
        "설비코드, 공구번호, 사용횟수",
        50,
        "오기술",
        "개발",
        None,
    ],
    [
        "IF-HB-011",
        "에너지 사용량",
        "SCADA",
        "ERP",
        "File",
        "일 1회",
        "설비코드, 전력량(kWh)",
        1,
        "강설비",
        "운영",
        None,
    ],
    [
        "IF-HB-012",
        "자재 투입 실적",
        "MES",
        "ERP",
        "API",
        "실시간",
        "LOT번호, 자재코드, 투입량",
        2500,
        "한생산",
        "운영",
        None,
    ],
]

TERMS = [
    TermIn(
        term="작업지시",
        definition="ERP 생산계획을 공정·설비 단위로 내린 생산 지시",
        abbreviation="WO",
        aliases=[AliasIn(alias="오더", department="생산관리팀"), AliasIn(alias="지시서", department="가공팀")],
    ),
    TermIn(
        term="라우팅",
        definition="품목을 만드는 공정 순서와 공정별 표준 시간",
        aliases=[AliasIn(alias="공정순서", department="생산기술팀")],
    ),
    TermIn(
        term="사이클타임",
        definition="한 개를 가공하는 데 걸리는 시간(초)",
        abbreviation="C/T",
        aliases=[AliasIn(alias="가공시간", department="가공팀"), AliasIn(alias="택트", department="생산관리팀")],
    ),
    TermIn(
        term="설비종합효율",
        definition="시간 가동률 \u00d7 성능 가동률 \u00d7 양품률",
        abbreviation="OEE",
        aliases=[AliasIn(alias="가동효율", department="설비보전팀")],
    ),
    TermIn(
        term="비가동",
        definition="계획 시간 중 설비가 생산하지 못한 시간(고장, 셋업, 자재 대기)",
        aliases=[AliasIn(alias="로스타임", department="가공팀"), AliasIn(alias="정지시간", department="설비보전팀")],
    ),
    TermIn(
        term="로트",
        definition="같은 소재·설비·교대에서 가공한 추적 단위",
        abbreviation="LOT",
        aliases=[AliasIn(alias="LOT번호", department="품질보증팀"), AliasIn(alias="배치", department="물류팀")],
    ),
    TermIn(
        term="초중종물 검사",
        definition="교대 시작·중간·종료 시점에 첫/중간/마지막 제품을 측정하는 검사",
        aliases=[AliasIn(alias="초품검사", department="가공팀")],
    ),
    TermIn(
        term="통계적 공정관리",
        definition="측정값의 관리도로 공정 이상을 조기에 찾는 방법",
        abbreviation="SPC",
        aliases=[AliasIn(alias="관리도", department="품질보증팀")],
    ),
    TermIn(
        term="부적합",
        definition="규격을 벗어난 제품 또는 그 처리 기록",
        abbreviation="NCR",
        aliases=[AliasIn(alias="불량", department="가공팀"), AliasIn(alias="NG", department="생산관리팀")],
    ),
    TermIn(
        term="재작업",
        definition="부적합품을 다시 가공해 규격 안으로 만드는 작업",
        aliases=[AliasIn(alias="리워크", department="가공팀")],
    ),
    TermIn(
        term="자재명세서",
        definition="품목을 이루는 자재와 소요량 목록",
        abbreviation="BOM",
        aliases=[AliasIn(alias="부품표", department="생산기술팀")],
    ),
    TermIn(
        term="공구 수명",
        definition="절삭 공구를 교체하기 전까지 허용하는 가공 횟수",
        aliases=[AliasIn(alias="툴 라이프", department="가공팀")],
    ),
    TermIn(
        term="셋업",
        definition="품목 전환 시 지그·공구·프로그램을 바꾸는 작업",
        aliases=[
            AliasIn(alias="기종 변경", department="가공팀"),
            AliasIn(alias="모델 체인지", department="생산관리팀"),
        ],
    ),
    TermIn(
        term="양품수량", definition="검사에 합격한 생산 수량", aliases=[AliasIn(alias="OK 수량", department="가공팀")]
    ),
    TermIn(
        term="교대", definition="주간·야간 등 작업 근무조", aliases=[AliasIn(alias="근무조", department="생산관리팀")]
    ),
]

INTERVIEWS = [
    (
        SubjectIn(engagement_id="", name="한생산", department="생산관리팀", job_title="팀장"),
        "생산관리팀 인터뷰",
        "실적 집계가 엑셀 수작업이라 ERP 반영이 하루 늦고, 비가동 원인이 기록되지 않는다.",
        [
            ("작업지시는 어떻게 내려가나요?", "ERP 계획을 엑셀로 받아 작업지시서를 출력해 현장에 배포합니다."),
            ("생산실적은 언제 ERP에 반영되나요?", "작업일보를 다음 날 오전에 엑셀로 모아 ERP에 수기 입력합니다."),
            ("가장 시간이 많이 드는 일은?", "교대마다 일보 숫자를 맞추는 데 하루 1~2시간이 듭니다."),
        ],
        [
            ("실적이 ERP에 하루 늦게 반영되어 납기 판단이 늦다", ["실적", "지연"], "MES→ERP 실적 자동 송신 범위 정의"),
            (
                "비가동 사유가 일보에 남지 않아 OEE를 계산할 수 없다",
                ["비가동", "OEE"],
                "SCADA 가동 신호와 비가동 사유 코드 설계",
            ),
        ],
    ),
    (
        SubjectIn(engagement_id="", name="윤품질", department="품질보증팀", job_title="과장"),
        "품질보증팀 인터뷰",
        "치수 측정값이 종이 성적서에만 있어 LOT 추적과 SPC가 불가능하다.",
        [
            ("검사 결과는 어디에 기록하나요?", "초중종물 측정값을 종이 성적서에 적고 월말에 스캔합니다."),
            ("고객 클레임 시 추적은 어떻게 하나요?", "LOT 번호로 성적서 파일을 찾는 데 반나절이 걸립니다."),
        ],
        [
            ("측정값이 전산화되지 않아 SPC 관리도를 그릴 수 없다", ["품질", "SPC"], "측정 항목·관리한계 마스터 정리"),
            ("LOT 추적에 반나절이 걸린다", ["LOT", "추적"], "LOT-설비-작업자 연결 키 정의"),
        ],
    ),
]

WORKBOOK = {
    "품목": (
        ["품목코드", "품목명", "규격", "재질", "표준사이클타임"],
        [
            ["HB-SH-100", "드라이브 샤프트", "\u00d832\u00d7410", "SCM440", 95],
            ["HB-HB-210", "휠 허브", "\u00d8120", "S45C", 140],
            ["HB-KN-305", "스티어링 너클", "LH", "FCD450", 180],
            ["HB-BR-410", "브레이크 캘리퍼 브래킷", "RH", "FCD500", 120],
        ],
    ),
    "설비": (
        ["설비코드", "설비명", "공정", "제조사", "도입일"],
        [
            ["CNC-01", "CNC 선반 1호", "선삭", "두산", date(2019, 3, 4)],
            ["CNC-02", "CNC 선반 2호", "선삭", "두산", date(2020, 6, 1)],
            ["MCT-01", "머시닝센터 1호", "밀링", "화천", date(2018, 11, 20)],
            ["GRD-01", "원통연삭기", "연삭", "현대위아", date(2021, 2, 15)],
        ],
    ),
    "작업실적": (
        ["실적번호", "작업일자", "교대", "품목코드", "설비코드", "생산수량", "불량수량", "작업자"],
        [
            ["R-2409-0001", date(2024, 9, 2), "주간", "HB-SH-100", "CNC-01", 310, 4, "김가공"],
            ["R-2409-0002", date(2024, 9, 2), "야간", "HB-SH-100", "CNC-01", 285, 7, "이가공"],
            ["R-2409-0003", date(2024, 9, 2), "주간", "HB-HB-210", "MCT-01", 160, 2, "박가공"],
            ["R-2409-0004", date(2024, 9, 3), "주간", "HB-KN-305", "MCT-01", 120, 5, "박가공"],
            ["R-2409-0005", date(2024, 9, 3), "주간", "HB-BR-410", "CNC-02", 240, 1, "최가공"],
            ["R-2409-0006", date(2024, 9, 3), "야간", "HB-SH-100", "GRD-01", 300, 3, "정가공"],
        ],
    ),
}


def _workbook() -> bytes:
    wb = Workbook()
    first = True
    for title, (header, rows) in WORKBOOK.items():
        ws = wb.active if first and wb.active is not None else wb.create_sheet(title)
        ws.title = title
        first = False
        ws.append(header)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _graph(lanes: list[str], steps: list[tuple[str, str, str, str | None]], systems: dict[str, str]) -> FlowGraph:
    """`steps` are (type, label, lane, system short name) chained in order."""
    nodes = [
        GraphNode(
            id=f"n{i}",
            type=kind,
            label=label,
            lane=lane,
            system_id=systems[short] if short else None,
            position=Position(x=60 + i * 190, y=40 + lanes.index(lane) * 150),
        )
        for i, (kind, label, lane, short) in enumerate(steps)
    ]
    edges = [GraphEdge(id=f"e{i}", source=f"n{i}", target=f"n{i + 1}") for i in range(len(nodes) - 1)]
    return FlowGraph(lanes=lanes, nodes=nodes, edges=edges)


AS_IS_LANES = ["생산관리", "가공 현장", "품질", "ERP"]
AS_IS = [
    ("start", "생산계획 확정", "생산관리", None),
    ("task", "ERP 계획 엑셀 다운로드", "생산관리", "ERP"),
    ("document", "작업지시서", "생산관리", None),
    ("task", "CNC 가공", "가공 현장", None),
    ("document", "작업일보", "가공 현장", None),
    ("decision", "초중종물 검사 합격?", "품질", None),
    ("document", "검사성적서", "품질", None),
    ("task", "실적 엑셀 집계", "생산관리", None),
    ("system", "ERP 실적 수기 입력", "ERP", "ERP"),
    ("end", "실적 마감", "ERP", None),
]
TO_BE_LANES = ["생산관리", "가공 현장", "품질", "시스템"]
TO_BE = [
    ("start", "생산계획 확정", "생산관리", None),
    ("system", "작업지시 I/F 수신", "시스템", "MES"),
    ("document", "전자 작업지시", "가공 현장", None),
    ("task", "태블릿 작업 시작", "가공 현장", None),
    ("system", "설비 가동 자동 수집", "시스템", "SCADA"),
    ("system", "측정값 SPC 자동 판정", "품질", "QMS"),
    ("document", "LOT 이력", "품질", None),
    ("system", "실적 자동 송신", "시스템", "ERP"),
    ("end", "실시간 실적 마감", "생산관리", None),
]

Attr = tuple[str, DataType, str | None, str]
CONCEPTS: list[tuple[str, str, str, list[Attr], str | None]] = [
    # name, upper key, definition, attributes (name, type, unit, column), table
    (
        "품목",
        "product",
        "고객에게 납품하는 가공 부품",
        [
            ("품목코드", "code", None, "품목코드"),
            ("규격", "string", None, "규격"),
            ("표준사이클타임", "integer", "초", "표준사이클타임"),
        ],
        "품목",
    ),
    (
        "설비",
        "equipment",
        "가공에 쓰는 CNC·MCT·연삭기",
        [("설비코드", "code", None, "설비코드"), ("공정", "string", None, "공정")],
        "설비",
    ),
    (
        "작업실적",
        "work_order",
        "작업지시에 대한 교대별 생산 결과",
        [
            ("실적번호", "code", None, "실적번호"),
            ("생산수량", "integer", "EA", "생산수량"),
            ("불량수량", "integer", "EA", "불량수량"),
        ],
        "작업실적",
    ),
    ("로트", "lot", "같은 소재·설비·교대로 가공한 추적 단위", [], None),
    ("품질검사", "inspection", "초중종물 치수 검사와 SPC 판정", [], None),
]
RELATIONS: list[tuple[str, str, str, Cardinality, str, str | None]] = [
    ("작업실적", "생산한다", "품목", "N:M", "생산된다", "품목코드"),
    ("작업실적", "사용한다", "설비", "N:M", "사용된다", "설비코드"),
    ("품질검사", "검사한다", "로트", "1:N", "검사된다", None),
]


def _asset_ref(db: Session, key: str) -> AssetRef:
    asset_id = db.scalars(select(AssetItem.asset_id).where(AssetItem.asset_key == key)).first()
    if asset_id is None:
        raise SystemExit(f"seed asset {key!r} missing")
    item = latest(db, asset_id)
    return AssetRef(asset_id=item.asset_id, version=item.version)


def seed_mes_sample(db: Session, principal: Principal, password: str) -> str:
    """Creates the sample tenant through the module services; returns its id."""
    if db.scalars(select(Tenant).where(Tenant.code == CODE)).first():
        raise SystemExit("sample tenant already exists")
    tenant = create_tenant(
        principal,
        db,
        TenantIn(name="한빛정밀(가상)", code=CODE, notes="자동차 부품 CNC 가공사 MES 고도화 샘플(가상 데이터)"),
    )
    fde = create_user(
        principal,
        db,
        UserIn(email=FDE_EMAIL, name="한빛 담당 FDE", role="fde", password=password, tenant_ids=[tenant.id]),
    )
    create_user(
        principal,
        db,
        UserIn(
            email=CLIENT_EMAIL,
            name="한빛 고객 관리자",
            role="client_admin",
            password=password,
            home_tenant_id=tenant.id,
        ),
    )
    ctx = TenantContext(
        principal=Principal(user=fde, ip=None, allowed_tenant_ids=frozenset({tenant.id})), tenant_id=tenant.id
    )
    with tenant_scope(db, tenant.id):
        _fill(db, ctx, fde.id)
    return tenant.id


def _fill(db: Session, ctx: TenantContext, fde_id: str) -> None:
    eng = create_engagement(
        ctx,
        db,
        EngagementIn(
            name="MES 고도화 진단",
            status="in_progress",
            start_date=date(2024, 9, 2),
            lead_fde_id=fde_id,
            description="수기 작업일보·종이 성적서를 MES·QMS로 전환하고 ERP 실적 반영을 실시간화한다.",
        ),
    )
    systems: dict[str, System] = {}
    for body in SYSTEMS:
        system = create_system(ctx, db, body)
        systems[body.short_name or body.name] = system
    sys_ids = {k: s.id for k, s in systems.items()}

    upload = upload_workbook(
        ctx, db, UploadFile(io.BytesIO(build_workbook(INTERFACES, [])), filename="hanbit-interfaces.xlsx")
    )
    apply_upload(ctx, db, upload, ApplyIn())
    for term in TERMS:
        create_term(ctx, db, term)

    session_ids: list[str] = []
    for subject_in, title, summary, qas, insights in INTERVIEWS:
        subject = create_subject(
            ctx, db, subject_in.model_copy(update={"engagement_id": eng.id, "system_ids": [sys_ids["MES"]]})
        )
        session = create_session(
            ctx,
            db,
            SessionIn(
                engagement_id=eng.id,
                subject_id=subject.id,
                title=title,
                session_date=date(2024, 9, 4),
                status="done",
                summary=summary,
            ),
        )
        session_ids.append(session.id)
        questions = [add_question(ctx, db, session, SessionQuestionIn(custom_text=q, answer=a)) for q, a in qas]
        for i, (text, tags, action) in enumerate(insights):
            insight = create_insight(
                ctx, db, session, InsightIn(session_question_id=questions[i].id, text=text, tags=tags)
            )
            create_action_item(
                ctx,
                db,
                session,
                ActionItemIn(insight_id=insight.id, title=action, assignee=subject.name, due=date(2024, 9, 20)),
            )

    as_is = create_flow(
        ctx,
        db,
        FlowIn(
            engagement_id=eng.id,
            kind="as_is",
            title="작업지시~실적 마감 (현행)",
            graph=_graph(AS_IS_LANES, AS_IS, sys_ids),
            description="엑셀·종이 기반 현행 흐름",
        ),
    )
    create_snapshot(ctx, db, as_is, SnapshotIn(note="현장 인터뷰 반영본"))
    to_be = create_flow(
        ctx,
        db,
        FlowIn(
            engagement_id=eng.id,
            kind="to_be",
            title="작업지시~실적 마감 (목표)",
            graph=_graph(TO_BE_LANES, TO_BE, sys_ids),
            description="MES·QMS·SCADA 연동 목표 흐름",
        ),
    )
    pair_flow(ctx, db, as_is, PairIn(flow_id=to_be.id))

    analysis = create_analysis(ctx, db, eng.id, UploadFile(io.BytesIO(_workbook()), filename="hanbit-production.xlsx"))
    confirm_erd(ctx, db, draft_of(db, ctx, analysis))

    upper = _asset_ref(db, "manufacturing-common")
    concepts = {}
    attrs = {}
    for name, key, definition, attributes, table in CONCEPTS:
        concept = create_concept(
            ctx,
            db,
            ConceptIn(
                name=name,
                status="confirmed",
                definition=definition,
                parent_ref=UpperRef(asset_id=upper.asset_id, version=upper.version, concept_key=key),
            ),
        )
        concepts[name] = concept
        create_mapping(ctx, db, MappingIn(target_kind="concept", target_id=concept.id, system_id=sys_ids["MES"]))
        for attr_name, data_type, unit, column in attributes:
            attr = create_attribute(
                ctx,
                db,
                concept,
                AttributeIn(name=attr_name, data_type=data_type, unit=unit, required=True),
            )
            attrs[(name, attr_name)] = attr
            create_mapping(
                ctx,
                db,
                MappingIn(
                    target_kind="attribute",
                    target_id=attr.id,
                    system_id=sys_ids["MES"],
                    table_name=table,
                    column_name=column,
                    origin="exmigrate",
                ),
            )
    for src, rel_name, dst, cardinality, inverse, fk_column in RELATIONS:
        rel = create_relation(
            ctx,
            db,
            RelationIn(
                source_concept_id=concepts[src].id,
                name=rel_name,
                target_concept_id=concepts[dst].id,
                cardinality=cardinality,
                inverse_name=inverse,
            ),
        )
        if fk_column:
            create_mapping(
                ctx,
                db,
                MappingIn(
                    target_kind="relation",
                    target_id=rel.id,
                    system_id=sys_ids["MES"],
                    table_name="작업실적",
                    column_name=fk_column,
                    origin="exmigrate",
                ),
            )
    for source in ("exmigrate_erd", "interface", "flowdesk_flow"):
        run_extract(ctx, db, ExtractIn(source_type=source))

    spec = create_document(
        ctx,
        db,
        DocumentIn(
            engagement_id=eng.id,
            doc_type="spec",
            title="한빛정밀 MES 실적 자동화 Spec",
            rule_pack_refs=[_asset_ref(db, "default-rules")],
            sources=SpecSources(
                discovery_session_ids=session_ids,
                flow_ids=[as_is.id, to_be.id],
                erd_analysis_ids=[analysis.id],
                interfaces=True,
                glossary=True,
                requirements=(
                    "- 교대 종료 10분 안에 실적이 ERP에 반영된다.\n- LOT 번호로 설비·작업자·측정값을 1분 안에 조회한다."
                ),
            ),
        ),
    )
    update_document(ctx, db, spec, DocumentPatch(status="confirmed"))
    create_document(
        ctx,
        db,
        DocumentIn(
            engagement_id=eng.id,
            doc_type="devin",
            title="한빛정밀 MES 실적 자동화 Devin.md",
            sources=SpecSources(flow_ids=[to_be.id], erd_analysis_ids=[analysis.id], interfaces=True),
        ),
    )

    project = create_project(
        ctx,
        db,
        ProjectIn(
            engagement_id=eng.id,
            name="MES 실적 자동 집계",
            status="active",
            stack="Python, FastAPI, MSSQL",
            description="작업실적·비가동·측정값을 MES에서 자동 집계해 ERP·QMS로 보낸다.",
            spec_document_ids=[spec.id],
        ),
    )
    for title, status, priority in (
        ("작업실적 엑셀 구조 분석", "done", "high"),
        ("ERP 실적 I/F(IF-HB-002) 매핑 정의", "in_progress", "high"),
        ("비가동 사유 코드 표준안", "review", "medium"),
        ("SPC 관리한계 마스터 이관", "todo", "medium"),
        ("현장 태블릿 화면 시안", "todo", "low"),
    ):
        create_task(ctx, db, project, TaskIn(title=title, status=status, priority=priority, assignee_id=fde_id))
    create_issue(
        ctx,
        db,
        project,
        IssueIn(
            title="야간 교대 실적이 다음 날짜로 집계됨",
            kind="bug",
            priority="high",
            description="22시~06시 실적의 작업일자 기준을 교대 시작일로 맞춰야 한다.",
        ),
    )
    create_issue(ctx, db, project, IssueIn(title="공구 교체 이력 수집 주기 확인", kind="question", priority="low"))
    create_instance(
        ctx,
        db,
        InstanceIn(
            name="한빛 생산 일보 요약",
            template_ref=_asset_ref(db, "daily-production-report"),
            engagement_id=eng.id,
            system_ids=[sys_ids["MES"], sys_ids["QMS"]],
            status="pilot",
            owner_id=fde_id,
            dev_project_id=project.id,
        ),
    )
