"""Workbook layout of the LUDA glossary template (seed asset ``glossary-excel``) and its reader/writer."""

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

SHEET = "용어 사전"
COLUMNS = (
    ("term", "표준 용어"),
    ("definition", "정의"),
    ("aliases", "부서별 호칭"),
    ("abbreviation", "약어"),
    ("concept", "관련 개념"),
    ("notes", "비고"),
)
STATUS_COLUMN = ("status", "상태")
REQUIRED = ("term",)
MAX_ROWS = 5000

STATUS_LABELS = {"candidate": "후보", "confirmed": "확정", "deprecated": "폐기"}
_STATUS_ALIASES = {
    **{v: k for k, v in STATUS_LABELS.items()},
    **{k: k for k in STATUS_LABELS},
}

_ALIAS_SPLIT = re.compile(r"[;\n,]+")
_ALIAS_WITH_PAREN = re.compile(r"^(?P<alias>.+?)\s*\((?P<dept>[^()]+)\)$")


class WorkbookError(Exception):
    def __init__(self, code: str, detail: dict[str, Any] | None = None) -> None:
        super().__init__(code)
        self.code = code
        self.detail = detail


@dataclass
class SheetRow:
    row: int
    values: dict[str, str | None] = field(default_factory=dict)


def parse_status(value: str | None) -> str | None:
    if value is None:
        return None
    return _STATUS_ALIASES.get(value.strip().lower(), _STATUS_ALIASES.get(value.strip()))


def parse_aliases(value: str | None) -> list[tuple[str, str | None]]:
    """``생산팀: 배합표; 품질팀: 레시피`` or ``배합표(생산팀)`` → [(alias, department)]."""
    out: list[tuple[str, str | None]] = []
    for raw in _ALIAS_SPLIT.split(value or ""):
        item = " ".join(raw.split())
        if not item:
            continue
        dept: str | None = None
        alias = item
        sep = next((s for s in (":", "\uff1a") if s in item), None)
        if sep is not None:
            left, right = item.split(sep, 1)
            dept, alias = left.strip() or None, right.strip()
        elif m := _ALIAS_WITH_PAREN.match(item):
            alias, dept = m["alias"].strip(), m["dept"].strip() or None
        if alias:
            out.append((alias, dept))
    return out


def format_aliases(aliases: list[tuple[str, str | None]]) -> str | None:
    text = "; ".join(f"{d}: {a}" if d else a for a, d in aliases)
    return text or None


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _read_sheet(rows: list[tuple[object, ...]]) -> list[SheetRow]:
    labels = dict(COLUMNS)
    if not rows:
        raise WorkbookError("invalid_workbook", {"sheet": SHEET, "missing_columns": [labels[k] for k in REQUIRED]})
    header = [_cell_text(v) or "" for v in rows[0]]
    positions = {key: header.index(label) for key, label in (*COLUMNS, STATUS_COLUMN) if label in header}
    missing = [labels[k] for k in REQUIRED if k not in positions]
    if missing:
        raise WorkbookError("invalid_workbook", {"sheet": SHEET, "missing_columns": missing})
    out: list[SheetRow] = []
    for idx, raw in enumerate(rows[1:], start=2):
        values = {k: _cell_text(raw[p]) if p < len(raw) else None for k, p in positions.items()}
        if any(v is not None for v in values.values()):
            out.append(SheetRow(row=idx, values=values))
        if len(out) > MAX_ROWS:
            raise WorkbookError("too_many_rows", {"sheet": SHEET, "max": MAX_ROWS})
    return out


def parse_workbook(data: bytes) -> list[SheetRow]:
    """Read values only (no formulas or macros are evaluated)."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise WorkbookError("invalid_workbook") from exc
    try:
        if SHEET not in wb.sheetnames:
            raise WorkbookError("invalid_workbook", {"missing_sheet": SHEET})
        return _read_sheet(list(wb[SHEET].iter_rows(values_only=True)))
    finally:
        wb.close()


_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")


def build_workbook(rows: list[list[Any]], with_status: bool) -> bytes:
    headers = [h for _, h in COLUMNS] + ([STATUS_COLUMN[1]] if with_status else [])
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = SHEET
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"
    for i, h in enumerate(headers, start=1):
        wide = h in ("정의", "부서별 호칭")
        ws.column_dimensions[get_column_letter(i)].width = 40 if wide else max(12, len(h) * 2 + 4)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
