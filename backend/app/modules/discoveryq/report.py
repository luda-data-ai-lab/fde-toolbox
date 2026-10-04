"""Offline Word (.docx) reports for DiscoveryQ sessions and engagements."""

import re
from collections import Counter
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.styles.style import ParagraphStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.tenancy.context import TenantContext
from app.core.tenancy.repository import TenantScopedRepository
from app.core.tenants.models import Engagement
from app.db.base import Base, utcnow
from app.modules.discoveryq.models import DiscoverySession
from app.modules.discoveryq.schemas import Worksheet
from app.modules.discoveryq.service import LABELS, Lang, _record_export, worksheet

FONT = "Malgun Gothic"
SOURCE_TYPE = "discovery_session"
TERM_STATUSES = ("open", "accepted", "merged")
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")

REPORT_LABELS: dict[Lang, dict[str, str]] = {
    "ko": {
        "session_report": "DiscoveryQ 세션 보고서",
        "engagement_report": "DiscoveryQ 진단 보고서",
        "engagement": "인게이지먼트",
        "period": "기간",
        "generated": "생성 일시",
        "overview": "개요",
        "session_count": "세션 수",
        "subject_count": "인터뷰 대상자 수",
        "insight_count": "인사이트 수",
        "action_count": "액션 아이템 수",
        "term_count": "등록 용어 수",
        "session_list": "세션 목록",
        "top_tags": "주요 태그",
        "session_details": "세션별 기록",
        "other_insights": "기타 인사이트",
        "tags": "태그",
        "terms": "등록 용어",
        "term": "용어",
        "definition": "정의",
        "aliases": "부서별 호칭",
        "term_open": "후보",
        "term_accepted": "확정",
        "term_merged": "병합",
        "none": "없음",
    },
    "en": {
        "session_report": "DiscoveryQ session report",
        "engagement_report": "DiscoveryQ discovery report",
        "engagement": "Engagement",
        "period": "Period",
        "generated": "Generated",
        "overview": "Overview",
        "session_count": "Sessions",
        "subject_count": "Interviewees",
        "insight_count": "Insights",
        "action_count": "Action items",
        "term_count": "Registered terms",
        "session_list": "Sessions",
        "top_tags": "Top tags",
        "session_details": "Session records",
        "other_insights": "Other insights",
        "tags": "Tags",
        "terms": "Registered terms",
        "term": "Term",
        "definition": "Definition",
        "aliases": "Department aliases",
        "term_open": "Candidate",
        "term_accepted": "Confirmed",
        "term_merged": "Merged",
        "none": "None",
    },
}


@dataclass
class TermRow:
    session_id: str | None
    term_id: str | None
    term: str
    definition: str
    aliases: str
    status: str


def _clean(value: object) -> str:
    return "" if value is None else _INVALID_XML.sub("", str(value))


def _new_document(title: str) -> DocxDocument:
    doc = Document()
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        assert isinstance(style, ParagraphStyle)
        style.font.name = FONT
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    doc.core_properties.title = title
    doc.core_properties.author = "LUDA FDE Toolbox"
    doc.add_heading(_clean(title), level=0)
    return doc


def _table(doc: DocxDocument, header: Sequence[str], rows: Iterable[Sequence[object]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = _clean(text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        for cell, value in zip(table.add_row().cells, row, strict=True):
            cell.text = _clean(value)


def _key_values(doc: DocxDocument, pairs: Sequence[tuple[str, object]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in pairs:
        cells = table.add_row().cells
        cells[0].text = _clean(key)
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = _clean(value)


def _to_bytes(doc: DocxDocument) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _terms(ctx: TenantContext, db: Session, session_ids: Sequence[str], lang: Lang) -> list[TermRow]:
    """Glossary terms registered from the given sessions.

    OntoMap tables are read by name (modules must not import each other) with an explicit tenant filter.
    """
    if not session_ids:
        return []
    rl = REPORT_LABELS[lang]
    tables = Base.metadata.tables
    cand, term, alias = tables["onto_candidates"], tables["onto_terms"], tables["onto_term_aliases"]
    candidates = db.execute(
        select(cand.c.name, cand.c.status, cand.c.source_id, cand.c.resolved_into_id)
        .where(
            cand.c.tenant_id == ctx.tenant_id,
            cand.c.source_type == SOURCE_TYPE,
            cand.c.source_id.in_(session_ids),
            cand.c.status.in_(TERM_STATUSES),
        )
        .order_by(cand.c.created_at)
    ).all()
    term_ids = [c.resolved_into_id for c in candidates if c.resolved_into_id]
    terms: dict[str, tuple[str, str | None]] = {}
    aliases: dict[str, list[str]] = {}
    if term_ids:
        for row in db.execute(
            select(term.c.id, term.c.term, term.c.definition).where(
                term.c.tenant_id == ctx.tenant_id, term.c.id.in_(term_ids)
            )
        ):
            terms[row.id] = (row.term, row.definition)
        for row in db.execute(
            select(alias.c.term_id, alias.c.alias, alias.c.department)
            .where(alias.c.tenant_id == ctx.tenant_id, alias.c.term_id.in_(term_ids))
            .order_by(alias.c.created_at)
        ):
            label = f"{row.department}: {row.alias}" if row.department else row.alias
            aliases.setdefault(row.term_id, []).append(label)
    rows = []
    for c in candidates:
        term_id = c.resolved_into_id if c.resolved_into_id in terms else None
        name, definition = terms[term_id] if term_id else (c.name, None)
        rows.append(
            TermRow(
                session_id=c.source_id,
                term_id=term_id,
                term=name,
                definition=definition or "",
                aliases=", ".join(aliases.get(term_id, [])) if term_id else "",
                status=rl[f"term_{c.status}"],
            )
        )
    return rows


def _terms_table(doc: DocxDocument, rows: Sequence[TermRow], lang: Lang) -> None:
    rl, lb = REPORT_LABELS[lang], LABELS[lang]
    _table(
        doc,
        [rl["term"], rl["definition"], rl["aliases"], lb["status"]],
        [[r.term, r.definition, r.aliases, r.status] for r in rows],
    )


def _session_meta(ws: Worksheet, lang: Lang) -> list[tuple[str, object]]:
    lb = LABELS[lang]
    s = ws.session
    pairs: list[tuple[str, object]] = [(lb["type"], lb.get(s.type, s.type))]
    if s.session_date:
        pairs.append((lb["date"], s.session_date.isoformat()))
    if ws.subject:
        extra = ", ".join(x for x in (ws.subject.department, ws.subject.job_title) if x)
        pairs.append((lb["subject"], ws.subject.name + (f" ({extra})" if extra else "")))
    pairs.append((lb["status"], lb.get(s.status, s.status)))
    return pairs


def _tagged(text: str, tags: Sequence[str]) -> str:
    return text + (" " + " ".join(f"#{t}" for t in tags) if tags else "")


def _write_session(doc: DocxDocument, ws: Worksheet, terms: Sequence[TermRow], lang: Lang, level: int) -> None:
    lb, rl = LABELS[lang], REPORT_LABELS[lang]
    _key_values(doc, _session_meta(ws, lang))
    if ws.session.summary:
        doc.add_heading(lb["summary"], level=level)
        doc.add_paragraph(_clean(ws.session.summary))

    doc.add_heading(lb["questions"], level=level)
    if not ws.questions:
        doc.add_paragraph(rl["none"])
    question_ids = {q.id for q in ws.questions}
    for n, q in enumerate(ws.questions, start=1):
        doc.add_heading(_clean(f"{n}. {q.text}"), level=min(level + 1, 9))
        if q.answer:
            doc.add_paragraph(_clean(q.answer))
        else:
            doc.add_paragraph().add_run(lb["no_answer"]).italic = True
        if q.follow_ups:
            p = doc.add_paragraph()
            p.add_run(f"{lb['follow_ups']}: ").bold = True
            p.add_run(_clean(" / ".join(q.follow_ups)))
        for insight in ws.insights:
            if insight.session_question_id == q.id:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{lb['insight']}: ").bold = True
                p.add_run(_clean(_tagged(insight.text, insight.tags)))

    others = [i for i in ws.insights if i.session_question_id not in question_ids]
    if others:
        doc.add_heading(rl["other_insights"], level=level)
        for insight in others:
            doc.add_paragraph(_clean(_tagged(insight.text, insight.tags)), style="List Bullet")

    doc.add_heading(lb["actions"], level=level)
    if ws.action_items:
        _table(
            doc,
            [lb["title"], lb["assignee"], lb["due"], lb["status"]],
            [
                [a.title, a.assignee, a.due.isoformat() if a.due else "", lb.get(a.status, a.status)]
                for a in ws.action_items
            ],
        )
    else:
        doc.add_paragraph(rl["none"])

    doc.add_heading(rl["terms"], level=level)
    if terms:
        _terms_table(doc, terms, lang)
    else:
        doc.add_paragraph(rl["none"])


def session_report(ctx: TenantContext, db: Session, session: DiscoverySession, lang: Lang) -> bytes:
    rl = REPORT_LABELS[lang]
    ws = worksheet(ctx, db, session)
    engagement = TenantScopedRepository(db, ctx, Engagement).get(session.engagement_id)
    doc = _new_document(f"{rl['session_report']}: {session.title}")
    pairs: list[tuple[str, object]] = [(rl["generated"], utcnow().strftime("%Y-%m-%d %H:%M UTC"))]
    if engagement:
        pairs.insert(0, (rl["engagement"], engagement.name))
    _key_values(doc, pairs)
    doc.add_paragraph()
    _write_session(doc, ws, _terms(ctx, db, [session.id], lang), lang, level=1)
    data = _to_bytes(doc)
    _record_export(ctx, db, "discoveryq.session_report", "discovery_session", session.id)
    return data


def engagement_report(ctx: TenantContext, db: Session, engagement: Engagement, lang: Lang) -> bytes:
    lb, rl = LABELS[lang], REPORT_LABELS[lang]
    sessions_repo = TenantScopedRepository(db, ctx, DiscoverySession)
    sessions = sessions_repo.all(
        sessions_repo.query()
        .where(DiscoverySession.engagement_id == engagement.id)
        .order_by(DiscoverySession.session_date, DiscoverySession.created_at)
    )
    sheets = [worksheet(ctx, db, s) for s in sessions]
    terms = _terms(ctx, db, [s.id for s in sessions], lang)
    insights = [i for ws in sheets for i in ws.insights]
    actions = [(ws.session.title, a) for ws in sheets for a in ws.action_items]
    unique_terms: dict[str, TermRow] = {}
    for row in terms:
        unique_terms.setdefault(row.term_id or f"candidate:{row.term}", row)

    doc = _new_document(f"{rl['engagement_report']}: {engagement.name}")
    period = " ~ ".join(d.isoformat() for d in (engagement.start_date, engagement.end_date) if d)
    _key_values(
        doc,
        [
            (rl["engagement"], engagement.name),
            *([(rl["period"], period)] if period else []),
            (rl["generated"], utcnow().strftime("%Y-%m-%d %H:%M UTC")),
        ],
    )
    if engagement.description:
        doc.add_paragraph(_clean(engagement.description))

    doc.add_heading(rl["overview"], level=1)
    _key_values(
        doc,
        [
            (rl["session_count"], len(sessions)),
            (rl["subject_count"], len({s.subject_id for s in sessions if s.subject_id})),
            (rl["insight_count"], len(insights)),
            (rl["action_count"], len(actions)),
            (rl["term_count"], len(unique_terms)),
        ],
    )
    tags = Counter(t for i in insights for t in i.tags)
    if tags:
        doc.add_heading(rl["top_tags"], level=2)
        doc.add_paragraph(_clean(", ".join(f"#{t} ({n})" for t, n in tags.most_common(10))))

    doc.add_heading(rl["session_list"], level=1)
    if sheets:
        _table(
            doc,
            [lb["session"], lb["date"], lb["type"], lb["subject"], lb["status"]],
            [
                [
                    ws.session.title,
                    ws.session.session_date.isoformat() if ws.session.session_date else "",
                    lb.get(ws.session.type, ws.session.type),
                    ws.subject.name if ws.subject else "",
                    lb.get(ws.session.status, ws.session.status),
                ]
                for ws in sheets
            ],
        )
    else:
        doc.add_paragraph(rl["none"])

    doc.add_heading(lb["actions"], level=1)
    if actions:
        _table(
            doc,
            [lb["session"], lb["title"], lb["assignee"], lb["due"], lb["status"]],
            [
                [title, a.title, a.assignee, a.due.isoformat() if a.due else "", lb.get(a.status, a.status)]
                for title, a in actions
            ],
        )
    else:
        doc.add_paragraph(rl["none"])

    doc.add_heading(rl["terms"], level=1)
    if unique_terms:
        _terms_table(doc, list(unique_terms.values()), lang)
    else:
        doc.add_paragraph(rl["none"])

    if sheets:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(rl["session_details"], level=1)
    for ws in sheets:
        doc.add_heading(_clean(ws.session.title), level=2)
        _write_session(doc, ws, [t for t in terms if t.session_id == ws.session.id], lang, level=3)

    data = _to_bytes(doc)
    _record_export(ctx, db, "discoveryq.engagement_report", "engagement", engagement.id)
    return data
