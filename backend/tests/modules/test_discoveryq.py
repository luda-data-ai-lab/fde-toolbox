import json
from io import BytesIO
from pathlib import Path
from typing import Any

import docx

from tests.conftest import World

SEED = Path(__file__).resolve().parents[2] / "seeds" / "assets" / "discoveryq-question-bank.json"


def _seed_bank(world: World) -> dict[str, Any]:
    pkg = json.loads(SEED.read_text(encoding="utf-8"))
    assert world.admin.post("/api/v1/assets/import", json=pkg).status_code == 200
    asset: dict[str, Any] = pkg["assets"][0]
    return asset


def _base(world: World) -> str:
    return f"/api/v1/t/{world.a.tenant_id}/discoveryq"


def test_question_bank_browse(world: World) -> None:
    base = _base(world)
    assert world.a.client_user.get(f"{base}/question-bank").json() == {"bank": None, "banks": []}
    asset = _seed_bank(world)
    body = world.a.client_user.get(f"{base}/question-bank").json()
    bank = body["bank"]
    assert bank["asset_id"] == asset["asset_id"] and bank["version"] == 1
    assert len(bank["categories"]) == 8 and all(len(c["questions"]) >= 8 for c in bank["categories"])
    assert bank["session_types"]["interview"] == ["status", "problem", "goal", "process"]
    assert body["banks"] == [{"asset_id": asset["asset_id"], "version": 1, "title": asset["title"]}]
    r = world.a.fde.get(f"{base}/question-bank", params={"asset_id": asset["asset_id"], "version": 9})
    assert r.status_code == 404


def test_subjects_and_sessions(world: World) -> None:
    base = _base(world)
    fde = world.a.fde
    eid = world.a.ids["engagement_id"]
    other_eng = fde.post(f"/api/v1/t/{world.a.tenant_id}/engagements", json={"name": "Other"}).json()["id"]
    r = fde.post(
        f"{base}/subjects", json={"engagement_id": eid, "name": "Lee", "system_ids": [world.b.ids["system_id"]]}
    )
    assert r.status_code == 422 and r.json()["error"]["detail"]["field"] == "system_ids"
    lee = fde.post(
        f"{base}/subjects", json={"engagement_id": eid, "name": "Lee", "department": "생산", "job_title": "팀장"}
    ).json()
    other_subject = fde.post(f"{base}/subjects", json={"engagement_id": other_eng, "name": "Park"}).json()
    r = fde.post(f"{base}/sessions", json={"engagement_id": eid, "title": "x", "subject_id": other_subject["id"]})
    assert r.status_code == 422 and r.json()["error"]["detail"]["field"] == "subject_id"
    s = fde.post(
        f"{base}/sessions",
        json={"engagement_id": eid, "title": "Coaching 1", "type": "coaching", "subject_id": lee["id"]},
    ).json()
    assert s["status"] == "planned" and s["type"] == "coaching"
    listed = fde.get(f"{base}/sessions", params={"type": "coaching"}).json()["items"]
    assert [x["id"] for x in listed] == [s["id"]]
    assert len(fde.get(f"{base}/sessions", params={"q": "Interview"}).json()["items"]) == 1
    assert [x["name"] for x in fde.get(f"{base}/subjects", params={"engagement_id": other_eng}).json()["items"]] == [
        "Park"
    ]
    patched = fde.patch(f"{base}/sessions/{s['id']}", json={"status": "done", "session_date": "2026-09-30"}).json()
    assert patched["status"] == "done" and patched["session_date"] == "2026-09-30"
    assert fde.delete(f"{base}/subjects/{lee['id']}").status_code == 204
    assert fde.get(f"{base}/sessions/{s['id']}").json()["session"]["subject_id"] is None
    assert world.a.client_admin.post(f"{base}/sessions", json={"engagement_id": eid, "title": "y"}).status_code == 403


def test_worksheet_flow(world: World) -> None:
    base = _base(world)
    fde = world.a.fde
    asset = _seed_bank(world)
    sid = world.a.ids["session_id"]
    ref = {"asset_id": asset["asset_id"], "version": 1, "question_id": "status-03"}
    q1 = fde.post(f"{base}/sessions/{sid}/questions", json={"question_ref": ref}).json()
    assert q1["text"] == "업무에 사용하는 시스템과 엑셀 파일은 무엇인가요?" and q1["category"] == "status"
    assert q1["follow_ups"] and q1["position"] == 1
    bad = {**ref, "question_id": "nope"}
    r = fde.post(f"{base}/sessions/{sid}/questions", json={"question_ref": bad})
    assert r.status_code == 422 and r.json()["error"]["detail"]["field"] == "question_ref"
    assert fde.post(f"{base}/sessions/{sid}/questions", json={}).status_code == 422
    both = {"custom_text": "a", "question_ref": ref}
    assert fde.post(f"{base}/sessions/{sid}/questions", json=both).status_code == 422
    q2 = fde.post(f"{base}/sessions/{sid}/questions", json={"custom_text": "MES 실적은 언제 마감하나요?"}).json()
    r = fde.post(f"{base}/sessions/{sid}/questions", json={"custom_question_id": world.b.ids["custom_question_id"]})
    assert r.status_code == 422
    fde.patch(f"{base}/session-questions/{q1['id']}", json={"answer": "SAP ERP와 '생산일보.xlsx'를 씁니다.\n매일 갱신"})
    ins = fde.post(
        f"{base}/sessions/{sid}/insights",
        json={"text": "생산일보 수기 집계", "tags": ["pain", " pain ", "", "data"], "session_question_id": q1["id"]},
    ).json()
    assert ins["tags"] == ["pain", "data"]
    other = fde.post(
        f"{base}/sessions", json={"engagement_id": world.a.ids["engagement_id"], "title": "Other session"}
    ).json()
    r = fde.post(f"{base}/sessions/{other['id']}/insights", json={"text": "x", "session_question_id": q1["id"]})
    assert r.status_code == 422
    r = fde.post(f"{base}/sessions/{other['id']}/action-items", json={"title": "x", "insight_id": ins["id"]})
    assert r.status_code == 422
    act = fde.post(
        f"{base}/sessions/{sid}/action-items",
        json={"title": "=HYPERLINK(1)", "assignee": "김대리", "due": "2026-10-15", "insight_id": ins["id"]},
    ).json()
    assert (
        fde.patch(f"{base}/action-items/{act['id']}", json={"status": "in_progress"}).json()["status"] == "in_progress"
    )
    ws = fde.get(f"{base}/sessions/{sid}").json()
    assert [q["id"] for q in ws["questions"]][1:] == [q1["id"], q2["id"]]
    assert ws["subject"]["name"] == "Kim TA" and len(ws["insights"]) == 2 and len(ws["action_items"]) == 2
    open_items = fde.get(f"{base}/action-items", params={"status": "in_progress"}).json()["items"]
    assert [a["id"] for a in open_items] == [act["id"]]
    by_eng = fde.get(f"{base}/action-items", params={"engagement_id": world.a.ids["engagement_id"]}).json()["items"]
    assert len(by_eng) == 2

    md = world.a.client_admin.get(f"{base}/sessions/{sid}/export.md")
    assert md.status_code == 200 and md.headers["content-type"].startswith("text/markdown")
    text = md.text
    assert text.startswith("# Interview TA\n")
    assert "- 대상자: Kim TA" in text and "### 2. 업무에 사용하는 시스템과 엑셀 파일은 무엇인가요?" in text
    assert "> SAP ERP와 '생산일보.xlsx'를 씁니다.\n> 매일 갱신" in text
    assert "`#pain` `#data`" in text and "| \\=HYPERLINK(1) |" not in text
    assert "| =HYPERLINK(1) | 김대리 | 2026-10-15 | 진행 중 |" in text
    en = world.a.fde.get(f"{base}/sessions/{sid}/export.md", params={"lang": "en"}).text
    assert "## Questions and answers" in en and "(no answer)" in en
    assert world.a.client_user.get(f"{base}/sessions/{sid}/export.md").status_code == 403

    csv_r = world.a.client_admin.get(f"{base}/action-items/export.csv")
    assert csv_r.status_code == 200
    lines = csv_r.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "세션,일자,항목,담당자,기한,상태,인사이트"
    assert "'=HYPERLINK(1)" in csv_r.text and "생산일보 수기 집계" in csv_r.text
    assert world.a.client_user.get(f"{base}/action-items/export.csv").status_code == 403
    actions = [
        a["action"]
        for a in world.a.fde.get(f"/api/v1/t/{world.a.tenant_id}/audit-logs", params={"limit": 200}).json()["items"]
    ]
    assert "discoveryq.session_export" in actions and "discoveryq.action_items_export" in actions

    assert fde.delete(f"{base}/session-questions/{q1['id']}").status_code == 204
    ws = fde.get(f"{base}/sessions/{sid}").json()
    assert next(i for i in ws["insights"] if i["id"] == ins["id"])["session_question_id"] is None
    assert fde.delete(f"{base}/sessions/{sid}").status_code == 204
    assert fde.get(f"{base}/action-items/{act['id']}").status_code == 404


def test_custom_questions(world: World) -> None:
    base = _base(world)
    fde = world.a.fde
    asset = _seed_bank(world)
    ref = {"asset_id": asset["asset_id"], "version": 1, "question_id": "goal-01"}
    clone = fde.post(
        f"{base}/custom-questions", json={"source_ref": ref, "engagement_id": world.a.ids["engagement_id"]}
    ).json()
    assert clone["category"] == "goal" and clone["source_ref"] == ref and clone["text"] and clone["follow_ups"]
    edited = fde.patch(f"{base}/custom-questions/{clone['id']}", json={"text": "도료 배합 목표는?"}).json()
    assert edited["text"] == "도료 배합 목표는?"
    assert fde.post(f"{base}/custom-questions", json={"category": "x"}).status_code == 422
    r = fde.post(f"{base}/custom-questions", json={"text": "x", "engagement_id": world.b.ids["engagement_id"]})
    assert r.status_code == 422
    other_eng = fde.post(f"/api/v1/t/{world.a.tenant_id}/engagements", json={"name": "Other"}).json()["id"]
    texts = {
        q["text"] for q in fde.get(f"{base}/custom-questions", params={"engagement_id": other_eng}).json()["items"]
    }
    assert texts == {"Custom question TA"}
    sq = fde.post(
        f"{base}/sessions/{world.a.ids['session_id']}/questions", json={"custom_question_id": clone["id"]}
    ).json()
    assert sq["text"] == "도료 배합 목표는?" and sq["category"] == "goal"
    assert fde.delete(f"{base}/custom-questions/{clone['id']}").status_code == 204
    ws = fde.get(f"{base}/sessions/{world.a.ids['session_id']}").json()
    kept = next(q for q in ws["questions"] if q["id"] == sq["id"])
    assert kept["custom_question_id"] is None and kept["text"] == "도료 배합 목표는?"
    assert world.a.client_user.get(f"{base}/custom-questions").status_code == 200
    assert world.a.client_user.post(f"{base}/custom-questions", json={"text": "x"}).status_code == 403


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_text(content: bytes) -> str:
    doc = docx.Document(BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def test_docx_reports(world: World) -> None:
    base = _base(world)
    fde = world.a.fde
    ids = world.a.ids
    tid = world.a.tenant_id
    answer = "배합비는 엑셀\x0b로 관리\n2줄"
    fde.patch(f"{base}/session-questions/{ids['session_question_id']}", json={"answer": answer})
    onto = f"/api/v1/t/{tid}/ontomap"
    cand = fde.post(
        f"{onto}/candidates",
        json={
            "name": "배합지시서",
            "source_type": "discovery_session",
            "source_id": ids["session_id"],
            "session_question_id": ids["session_question_id"],
        },
    ).json()
    accepted = fde.post(
        f"{onto}/candidates/{cand['id']}/accept",
        json={"definition": "배합 작업 지시 문서", "aliases": [{"alias": "배합표", "department": "생산팀"}]},
    )
    assert accepted.status_code == 200, accepted.text
    ignored = fde.post(
        f"{onto}/candidates",
        json={"name": "무시됨", "source_type": "discovery_session", "source_id": ids["session_id"]},
    ).json()
    assert fde.post(f"{onto}/candidates/{ignored['id']}/ignore").status_code == 200
    second = fde.post(
        f"{base}/sessions",
        json={"engagement_id": ids["engagement_id"], "title": "Coaching follow-up", "type": "coaching"},
    ).json()
    fde.post(f"{base}/sessions/{second['id']}/insights", json={"text": "Loose insight", "tags": ["품질"]})

    r = fde.get(f"{base}/sessions/{ids['session_id']}/report.docx")
    assert r.status_code == 200 and r.headers["content-type"] == DOCX
    assert f"discoveryq-session-{ids['session_id']}.docx" in r.headers["content-disposition"]
    text = _docx_text(r.content)
    for expected in (
        f"Interview {world.a.code}",
        f"Custom question {world.a.code}",
        "배합비는 엑셀로 관리\n2줄",
        f"Insight {world.a.code}",
        f"Action {world.a.code}",
        "배합지시서",
        "배합 작업 지시 문서",
        "생산팀: 배합표",
        "확정",
        f"Cand {world.a.code}",
        "후보",
    ):
        assert expected in text, expected
    assert "무시됨" not in text and "Coaching follow-up" not in text

    eng = fde.get(f"{base}/engagements/{ids['engagement_id']}/report.docx", params={"lang": "en"})
    assert eng.status_code == 200 and eng.headers["content-type"] == DOCX
    text = _docx_text(eng.content)
    for expected in (
        "DiscoveryQ discovery report",
        f"Interview {world.a.code}",
        "Coaching follow-up",
        "Loose insight #품질",
        "#품질 (1)",
        f"Action {world.a.code}",
        "생산팀: 배합표",
        "Confirmed",
    ):
        assert expected in text, expected
    assert f"Interview {world.b.code}" not in text

    assert world.a.client_user.get(f"{base}/sessions/{ids['session_id']}/report.docx").status_code == 403
    assert world.a.client_user.get(f"{base}/engagements/{ids['engagement_id']}/report.docx").status_code == 403
    assert fde.get(f"{base}/engagements/{world.b.ids['engagement_id']}/report.docx").status_code == 404
    actions = {e["action"] for e in fde.get(f"/api/v1/t/{tid}/audit-logs", params={"limit": 200}).json()["items"]}
    assert {"discoveryq.session_report", "discoveryq.engagement_report"} <= actions
