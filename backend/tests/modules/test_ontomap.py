import csv
import io
from typing import Any

from openpyxl import Workbook, load_workbook

from app.modules.ontomap import excel
from tests.conftest import World

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _base(world: World) -> str:
    return f"/api/v1/t/{world.a.tenant_id}/ontomap"


def _workbook(rows: list[list[Any]], headers: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = excel.SHEET
    ws.append(headers or [h for _, h in excel.COLUMNS] + ["상태"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _actions(world: World) -> list[str]:
    logs = world.a.fde.get(f"/api/v1/t/{world.a.tenant_id}/audit-logs", params={"limit": 200}).json()["items"]
    return [x["action"] for x in logs]


def test_parse_aliases() -> None:
    assert excel.parse_aliases("생산팀: 배합표; 품질팀\uff1a레시피\n처방(R&D)") == [
        ("배합표", "생산팀"),
        ("레시피", "품질팀"),
        ("처방", "R&D"),
    ]
    assert excel.parse_aliases("BOM, 자재명세") == [("BOM", None), ("자재명세", None)]
    assert excel.format_aliases([("배합표", "생산팀"), ("BOM", None)]) == "생산팀: 배합표; BOM"
    assert excel.parse_status("확정") == "confirmed" and excel.parse_status("Deprecated") == "deprecated"
    assert excel.parse_status("??") is None


def test_term_crud_and_aliases(world: World) -> None:
    base = _base(world)
    fde = world.a.fde
    body = {
        "term": "  레시피 ",
        "definition": "제품 배합 기준",
        "abbreviation": "RCP",
        "aliases": [
            {"alias": "배합표", "department": "생산팀"},
            {"alias": "배합표", "department": "생산팀"},
            {"alias": "처방", "department": "R&D"},
        ],
        "tenant_id": world.b.tenant_id,
    }
    r = fde.post(f"{base}/terms", json=body)
    assert r.status_code == 201, r.text
    term = r.json()
    assert term["term"] == "레시피" and term["tenant_id"] == world.a.tenant_id and term["status"] == "confirmed"
    assert term["aliases"] == [{"alias": "배합표", "department": "생산팀"}, {"alias": "처방", "department": "R&D"}]
    dup = fde.post(f"{base}/terms", json={"term": "레시피"})
    assert dup.status_code == 409 and dup.json()["error"]["detail"]["term_id"] == term["id"]
    assert fde.get(f"{base}/terms", params={"q": "배합"}).json()["items"][0]["id"] == term["id"]
    assert [x["id"] for x in fde.get(f"{base}/terms", params={"department": "R&D"}).json()["items"]] == [term["id"]]
    assert fde.get(f"{base}/terms", params={"status": "deprecated"}).json()["items"] == []
    patched = fde.patch(
        f"{base}/terms/{term['id']}", json={"status": "deprecated", "aliases": [{"alias": "Recipe"}]}
    ).json()
    assert patched["status"] == "deprecated" and patched["aliases"] == [{"alias": "Recipe", "department": None}]
    kept = fde.patch(f"{base}/terms/{term['id']}", json={"definition": "x"}).json()
    assert kept["aliases"] == patched["aliases"] and kept["term"] == "레시피"
    other = fde.post(f"{base}/terms", json={"term": "BOM"}).json()
    assert fde.patch(f"{base}/terms/{other['id']}", json={"term": "레시피"}).status_code == 409
    assert world.a.client_user.get(f"{base}/terms/{term['id']}").status_code == 200
    assert world.a.client_user.patch(f"{base}/terms/{term['id']}", json={"definition": "y"}).status_code == 403
    assert world.a.client_admin.post(f"{base}/terms", json={"term": "Z"}).status_code == 403
    assert fde.delete(f"{base}/terms/{term['id']}").status_code == 204
    assert fde.get(f"{base}/terms/{term['id']}").status_code == 404
    # foreign tenant's term id is invisible under own tenant path
    assert fde.get(f"{base}/terms/{world.b.ids['term_id']}").status_code == 404
    actions = _actions(world)
    for a in ("ontomap.term_create", "ontomap.term_update", "ontomap.term_delete"):
        assert a in actions


def test_candidates_from_coachq(world: World) -> None:
    base = _base(world)
    fde = world.a.fde
    sid = world.a.ids["session_id"]
    body = {
        "name": "배합 지시서",
        "source_type": "coach_session",
        "source_id": sid,
        "session_question_id": world.a.ids["session_question_id"],
        "context": "생산팀은 배합 지시서를 매일 출력한다",
        "department": "생산팀",
    }
    r = fde.post(f"{base}/candidates", json=body)
    assert r.status_code == 201, r.text
    cand = r.json()
    assert cand["status"] == "open" and cand["payload"]["department"] == "생산팀"
    again = fde.post(f"{base}/candidates", json=body)
    assert again.status_code == 200 and again.json()["id"] == cand["id"]
    for field, patch in (
        ("source_id", {"source_id": world.b.ids["session_id"]}),
        ("session_question_id", {"session_question_id": world.b.ids["session_question_id"]}),
        ("source_id", {"source_id": None}),
    ):
        bad = fde.post(f"{base}/candidates", json={**body, **patch})
        assert bad.status_code == 422 and bad.json()["error"]["detail"]["field"] == field
    listed = fde.get(f"{base}/candidates", params={"source_id": sid, "status": "open"}).json()["items"]
    assert {c["name"] for c in listed} == {"배합 지시서", "Cand TA"}

    # accept → confirmed term carrying source metadata and the department alias
    accepted = fde.post(f"{base}/candidates/{cand['id']}/accept", json={"definition": "배합 작업 지시 문서"}).json()
    assert accepted["status"] == "accepted"
    term = fde.get(f"{base}/terms/{accepted['resolved_into_id']}").json()
    assert term["status"] == "confirmed" and term["source_type"] == "coach_session" and term["source_id"] == sid
    assert term["definition"] == "배합 작업 지시 문서"
    assert term["aliases"] == [{"alias": "배합 지시서", "department": "생산팀"}]
    assert fde.post(f"{base}/candidates/{cand['id']}/ignore").status_code == 409

    # similar suggestions and merge into an existing term
    similar = fde.post(f"{base}/candidates", json={"name": "배합지시서 ", "source_type": "manual"}).json()
    assert similar["similar"][0] == {
        "term_id": term["id"],
        "term": "배합 지시서",
        "matched": "배합 지시서",
        "score": 100,
    }
    taken = fde.post(f"{base}/candidates/{similar['id']}/accept", json={"term": "배합 지시서"})
    assert taken.status_code == 409 and taken.json()["error"]["code"] == "term_taken"
    foreign = fde.post(f"{base}/candidates/{similar['id']}/merge", json={"term_id": world.b.ids["term_id"]})
    assert foreign.status_code == 422 and foreign.json()["error"]["detail"]["field"] == "term_id"
    merged = fde.post(
        f"{base}/candidates/{similar['id']}/merge", json={"term_id": term["id"], "department": "품질팀"}
    ).json()
    assert merged["status"] == "merged" and merged["resolved_into_id"] == term["id"]
    aliases = fde.get(f"{base}/terms/{term['id']}").json()["aliases"]
    assert {"alias": "배합지시서", "department": "품질팀"} in aliases

    # ignore / reopen
    noise = fde.post(f"{base}/candidates", json={"name": "그거"}).json()
    assert fde.post(f"{base}/candidates/{noise['id']}/ignore").json()["status"] == "ignored"
    assert fde.post(f"{base}/candidates/{noise['id']}/reopen").json()["status"] == "open"
    assert fde.post(f"{base}/candidates/{noise['id']}/reopen").status_code == 409
    assert world.a.client_user.post(f"{base}/candidates", json={"name": "x"}).status_code == 403
    assert fde.get(f"{base}/candidates/{world.b.ids['candidate_id']}").status_code == 404
    actions = _actions(world)
    for a in ("ontomap.candidate_create", "ontomap.candidate_accept", "ontomap.candidate_merge"):
        assert a in actions


def test_import_and_export(world: World) -> None:
    base = _base(world)
    fde = world.a.fde
    data = _workbook(
        [
            ["설비", "생산 장비", "생산팀: 라인; 보전팀: 머신", "EQ", "Equipment", None, None],
            ["Term TA", "updated def", "QA: TA2", None, None, None, "폐기"],
            [None, "no term", None, None, None, None, None],
            ["설비", "dup", None, None, None, None, None],
            ["=CMD()", "formula-like", None, None, None, None, "maybe"],
        ]
    )
    files = {"file": ("glossary.xlsx", data, XLSX)}
    preview = fde.post(f"{base}/terms/import", files=files)
    assert preview.status_code == 200, preview.text
    result = preview.json()
    assert result["applied"] is False
    assert result["summary"] == {"total": 5, "valid": 2, "invalid": 3, "create": 1, "update": 1}
    assert [e["code"] for e in result["rows"][2]["errors"]] == ["required"]
    assert [e["code"] for e in result["rows"][3]["errors"]] == ["duplicate_in_file"]
    # formulas are never evaluated: a formula cell without a cached value reads as empty
    assert {e["code"] for e in result["rows"][4]["errors"]} == {"invalid_status", "required"}
    assert fde.get(f"{base}/terms", params={"q": "설비"}).json()["items"] == []

    applied = fde.post(f"{base}/terms/import", params={"apply": "true"}, files=files).json()
    assert applied["applied"] is True
    eq = fde.get(f"{base}/terms", params={"q": "설비"}).json()["items"][0]
    assert eq["source_type"] == "import" and eq["abbreviation"] == "EQ"
    assert eq["aliases"] == [{"alias": "라인", "department": "생산팀"}, {"alias": "머신", "department": "보전팀"}]
    ta = fde.get(f"{base}/terms/{world.a.ids['term_id']}").json()
    assert ta["status"] == "deprecated" and ta["definition"] == "updated def"
    assert ta["aliases"] == [{"alias": "TTA", "department": "QA"}, {"alias": "TA2", "department": "QA"}]
    assert fde.get(f"{base}/terms/{world.b.ids['term_id']}").status_code == 404
    assert world.b.fde.get(f"/api/v1/t/{world.b.tenant_id}/ontomap/terms", params={"q": "설비"}).json()["items"] == []

    bad = fde.post(f"{base}/terms/import", files={"file": ("g.xlsx", _workbook([], headers=["x"]), XLSX)})
    assert bad.status_code == 400 and bad.json()["error"]["detail"]["missing_columns"] == ["표준 용어"]
    assert fde.post(f"{base}/terms/import", files={"file": ("g.txt", b"x", "text/plain")}).status_code == 400
    assert world.a.client_admin.post(f"{base}/terms/import", files=files).status_code == 403

    tpl = load_workbook(io.BytesIO(fde.get(f"{base}/terms/template.xlsx").content))
    assert [c.value for c in tpl[excel.SHEET][1]] == [h for _, h in excel.COLUMNS]

    r = fde.get(f"{base}/terms/export.xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    ws = load_workbook(io.BytesIO(r.content))[excel.SHEET]
    rows = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["설비"][2] == "생산팀: 라인; 보전팀: 머신" and rows["설비"][6] == "확정"
    assert "Term TB" not in rows
    only = load_workbook(io.BytesIO(fde.get(f"{base}/terms/export.xlsx", params={"status": "deprecated"}).content))
    assert [row[0] for row in only[excel.SHEET].iter_rows(min_row=2, values_only=True)] == ["Term TA"]
    text = fde.get(f"{base}/terms/export.csv").content.decode("utf-8-sig")
    assert next(csv.reader(io.StringIO(text)))[0] == "표준 용어"
    assert world.a.client_admin.get(f"{base}/terms/export.xlsx").status_code == 200
    assert world.a.client_user.get(f"{base}/terms/export.xlsx").status_code == 403
    actions = _actions(world)
    assert "ontomap.terms_import" in actions and actions.count("ontomap.terms_export") >= 3


def test_reimport_roundtrip(world: World) -> None:
    base = _base(world)
    fde = world.a.fde
    exported = fde.get(f"{base}/terms/export.xlsx").content
    r = fde.post(f"{base}/terms/import", files={"file": ("glossary.xlsx", exported, XLSX)})
    assert r.json()["summary"] == {"total": 1, "valid": 1, "invalid": 0, "create": 0, "update": 1}
    fde.post(f"{base}/terms/import", params={"apply": "true"}, files={"file": ("glossary.xlsx", exported, XLSX)})
    term = fde.get(f"{base}/terms/{world.a.ids['term_id']}").json()
    assert term["aliases"] == [{"alias": "TTA", "department": "QA"}]
