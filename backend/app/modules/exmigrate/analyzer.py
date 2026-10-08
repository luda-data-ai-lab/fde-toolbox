"""Workbook structure and formula analysis (pure: bytes in, report out; macros are never executed)."""

import io
import re
import zipfile
from collections import Counter
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import AppError
from app.modules.exmigrate.schemas import (
    ColumnReport,
    ColumnType,
    FormulaCell,
    FormulaSummary,
    FunctionUse,
    LookupRef,
    SheetRef,
    SheetReport,
    SheetWarning,
    WorkbookReport,
)

MAX_ROWS = 20_000
MAX_COLUMNS = 200
HEADER_SCAN_ROWS = 20
MAX_FORMULA_CELLS = 50
MAX_WARNINGS = 30

SIMPLE_FUNCTIONS = frozenset(
    [
        "SUM",
        "SUMIF",
        "SUMIFS",
        "COUNT",
        "COUNTA",
        "COUNTBLANK",
        "COUNTIF",
        "COUNTIFS",
        "AVERAGE",
        "AVERAGEIF",
        "AVERAGEIFS",
        "MIN",
        "MAX",
        "ROUND",
        "ROUNDUP",
        "ROUNDDOWN",
        "IF",
        "IFS",
        "IFERROR",
        "IFNA",
        "AND",
        "OR",
        "NOT",
        "VLOOKUP",
        "HLOOKUP",
        "XLOOKUP",
        "INDEX",
        "MATCH",
        "CONCATENATE",
        "CONCAT",
        "TEXT",
        "LEFT",
        "RIGHT",
        "MID",
        "LEN",
        "TRIM",
        "UPPER",
        "LOWER",
        "SUBSTITUTE",
        "TODAY",
        "NOW",
        "DATE",
        "YEAR",
        "MONTH",
        "DAY",
        "ABS",
        "INT",
        "MOD",
        "ISBLANK",
        "ISNUMBER",
        "VALUE",
    ]
)

_STRING = re.compile(r'"(?:[^"]|"")*"')
_FUNCTION = re.compile(r"\b([A-Z][A-Z0-9.]*)\s*\(")
_SHEET = r"(?:'((?:[^']|'')+)'|([^\s!'\",()=+\-*/&<>:;^%{}]+))"
_SHEET_REF = re.compile(_SHEET + r"!")
_LOOKUP = re.compile(
    r"\b(?:VLOOKUP|XLOOKUP|MATCH)\(\s*\$?([A-Z]{1,3})\$?\d+\s*,\s*" + _SHEET + r"!\$?([A-Z]{1,3})", re.IGNORECASE
)
_IDENTIFIER_HINT = re.compile(r"(^id$|_id$|^code$|_code$|코드|번호|^no$|_no$)", re.IGNORECASE)


def _formula_text(value: object) -> str | None:
    if isinstance(value, ArrayFormula):
        return value.text if isinstance(value.text, str) else None
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _value_type(value: object) -> ColumnType | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "integer" if value.is_integer() else "decimal"
    if isinstance(value, datetime):
        return "date" if value.time() == time(0) else "datetime"
    if isinstance(value, date):
        return "date"
    return "text"


def merge_types(types: set[ColumnType]) -> tuple[ColumnType | None, bool]:
    """Common type of a column's values and whether the values were mixed."""
    if not types:
        return None, False
    if len(types) == 1:
        return next(iter(types)), False
    if types <= {"integer", "decimal"}:
        return "decimal", False
    if types <= {"date", "datetime"}:
        return "datetime", False
    return "text", True


def _text(value: object, limit: int = 50) -> str:
    if isinstance(value, datetime) and value.time() == time(0):
        value = value.date()
    s = value.isoformat() if isinstance(value, date) else str(value)
    return s if len(s) <= limit else s[: limit - 1] + "…"


def _sheet_name(quoted: str | None, bare: str | None) -> str:
    return quoted.replace("''", "'") if quoted is not None else (bare or "")


def _header_row(ws: Worksheet, max_row: int, max_col: int) -> int | None:
    best, best_score = None, 0
    for r in range(1, min(max_row, HEADER_SCAN_ROWS) + 1):
        score = sum(
            1
            for c in range(1, max_col + 1)
            if isinstance(v := ws.cell(r, c).value, str) and v.strip() and _formula_text(v) is None
        )
        if score > best_score:
            best, best_score = r, score
    return best


class _SheetFormulas:
    def __init__(self, sheets: dict[str, str]) -> None:
        self.sheets = sheets
        self.count = 0
        self.complex = 0
        self.cells: list[FormulaCell] = []
        self.functions: Counter[str] = Counter()
        self.references: Counter[str] = Counter()
        self.lookups: Counter[tuple[str, str, str]] = Counter()
        self.external = 0

    def add(self, cell: str, formula: str) -> None:
        self.count += 1
        if len(self.cells) < MAX_FORMULA_CELLS:
            self.cells.append(FormulaCell(cell=cell, formula=formula[:300]))
        code = _STRING.sub('""', formula)
        names = {m.upper() for m in _FUNCTION.findall(code.upper())}
        self.functions.update(names)
        if any(n not in SIMPLE_FUNCTIONS for n in names):
            self.complex += 1
        for quoted, bare in _SHEET_REF.findall(code):
            name = _sheet_name(quoted or None, bare or None)
            if name.startswith("["):
                self.external += 1
            elif (target := self.sheets.get(name.lower())) is not None:
                self.references[target] += 1
        for column, quoted, bare, target_column in _LOOKUP.findall(code):
            target = self.sheets.get(_sheet_name(quoted or None, bare or None).lower())
            if target is not None:
                self.lookups[(column.upper(), target, target_column.upper())] += 1

    def summary(self) -> FormulaSummary:
        return FormulaSummary(
            count=self.count,
            complex=self.complex,
            cells=self.cells,
            functions=_functions(self.functions),
            references=[SheetRef(sheet=s, count=n) for s, n in self.references.most_common()],
            lookups=[
                LookupRef(column=c, target_sheet=s, target_column=tc, count=n)
                for (c, s, tc), n in self.lookups.most_common()
            ],
        )


def _functions(counter: Counter[str]) -> list[FunctionUse]:
    return [FunctionUse(name=n, count=c, simple=n in SIMPLE_FUNCTIONS) for n, c in counter.most_common()]


def _analyze_sheet(ws: Worksheet, values: Worksheet, sheets: dict[str, str]) -> SheetReport:
    max_row, max_col = ws.max_row or 0, ws.max_column or 0
    warnings: list[SheetWarning] = []

    def warn(code: str, where: str | None = None, detail: str | None = None) -> None:
        if len(warnings) < MAX_WARNINGS:
            warnings.append(SheetWarning(code=code, where=where, detail=detail))

    if max_col > MAX_COLUMNS:
        warn("truncated_columns", detail=str(max_col))
        max_col = MAX_COLUMNS
    if max_row > MAX_ROWS:
        warn("truncated_rows", detail=str(max_row))
        max_row = MAX_ROWS
    for rng in list(ws.merged_cells.ranges)[:10]:
        warn("merged_cells", where=str(rng))

    formulas = _SheetFormulas(sheets)
    grid: list[list[object]] = []
    formula_cols: Counter[int] = Counter()
    for r, (row, cached) in enumerate(
        zip(
            ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True),
            values.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True),
            strict=False,
        ),
        start=1,
    ):
        out: list[object] = []
        for c, (raw, val) in enumerate(zip(row, cached, strict=False), start=1):
            text = _formula_text(raw)
            if text is not None:
                formulas.add(f"{get_column_letter(c)}{r}", text)
                formula_cols[c] += 1
                out.append(val)
            else:
                out.append(raw)
        grid.append(out)
    if formulas.external:
        warn("external_reference", detail=str(formulas.external))

    header = _header_row(ws, max_row, max_col)
    start = (header or 0) + 1
    body = grid[start - 1 :]
    while body and all(_value_type(v) is None for v in body[-1]):
        body.pop()
    blank = [start + i for i, row in enumerate(body) if all(_value_type(v) is None for v in row)]
    if blank:
        warn("blank_rows", where=str(blank[0]), detail=str(len(blank)))
    if not body and header is None:
        warn("empty_sheet")

    columns: list[ColumnReport] = []
    seen: set[str] = set()
    for c in range(1, max_col + 1):
        letter = get_column_letter(c)
        raw_name = grid[header - 1][c - 1] if header else None
        name = str(raw_name).strip() if raw_name is not None and _value_type(raw_name) else ""
        cells = [row[c - 1] if c - 1 < len(row) else None for row in body]
        present = [v for v in cells if _value_type(v) is not None]
        if not name and not present:
            continue
        if not name and header:
            warn("empty_header", where=f"{letter}{header}")
        if name and name.lower() in seen:
            warn("duplicate_header", where=f"{letter}{header}", detail=name)
        seen.add(name.lower())
        kind, mixed = merge_types({t for v in present if (t := _value_type(v)) is not None})
        if mixed:
            warn("mixed_types", where=letter, detail=name or letter)
        distinct = list(dict.fromkeys(_text(v, 200) for v in present))
        columns.append(
            ColumnReport(
                letter=letter,
                name=name or letter,
                inferred_type=kind,
                null_ratio=round(1 - len(present) / len(body), 4) if body else 1.0,
                distinct=len(distinct),
                samples=[_text(v) for v in distinct[:3]],
                formula_cells=formula_cols[c],
            )
        )
    return SheetReport(
        name=ws.title,
        dimension=ws.dimensions,
        max_row=ws.max_row or 0,
        max_column=ws.max_column or 0,
        header_row=header,
        data_rows=len(body),
        columns=columns,
        warnings=warnings,
        formulas=formulas.summary(),
    )


def _has_macros(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return any(n.lower().endswith("vbaproject.bin") for n in zf.namelist())
    except zipfile.BadZipFile:
        return False


def analyze_workbook(data: bytes) -> WorkbookReport:
    try:
        formulas_wb = load_workbook(io.BytesIO(data), data_only=False, keep_vba=False)
        values_wb = load_workbook(io.BytesIO(data), data_only=True, keep_vba=False)
    except Exception as exc:  # openpyxl raises many unrelated types for corrupt input
        raise AppError(400, "workbook_invalid") from exc
    sheets = {ws.title.lower(): ws.title for ws in formulas_wb.worksheets}
    reports = [_analyze_sheet(ws, values_wb[ws.title], sheets) for ws in formulas_wb.worksheets]
    total: Counter[str] = Counter()
    for s in reports:
        total.update({f.name: f.count for f in s.formulas.functions})
    return WorkbookReport(
        sheets=reports,
        has_macros=_has_macros(data),
        formula_count=sum(s.formulas.count for s in reports),
        complex_formulas=sum(s.formulas.complex for s in reports),
        functions=_functions(total),
    )


def looks_like_identifier(name: str) -> bool:
    return bool(_IDENTIFIER_HINT.search(name.strip()))


def report_of(data: dict[str, Any]) -> WorkbookReport:
    return WorkbookReport.model_validate(data)
